import csv, json, os
from datetime import datetime, timedelta
from collections import defaultdict
from pathlib import Path

DATA_DIR = Path(__file__).parent.parent / "data-ref"

def read_csv(name):
    fp = DATA_DIR / name
    if not fp.exists(): return None
    with open(fp, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    return rows

def read_sku_master():
    rows = read_csv("sku_master.csv")
    if not rows: return {}, {}, {}, {}
    sku_attrs, sku_to_gb, sku_pallet, gb_style_color = {}, {}, {}, {}
    for r in rows:
        sku = r["SKU"]
        sku_attrs[sku] = {"Style": r.get("Style",""), "Color": r.get("Color",""),
                           "Usage": r.get("Usage", "")}
        sku_to_gb[sku] = r.get("GB_PN","")
        sku_pallet[sku] = int(r.get("Pallet_Qty", 864))
        gb = r.get("GB_PN","")
        if gb and r.get("Style") and r.get("Color"):
            gb_style_color[gb] = (r["Style"], r["Color"])
    return sku_attrs, sku_to_gb, sku_pallet, gb_style_color

def read_wide_csv(name, value_col=None):
    """Read wide CSV → {SKU: {date_str: value}}"""
    rows = read_csv(name)
    if not rows: return {}
    result = {}
    for r in rows:
        sku = r.get("SKU") or r.get("PN", "")
        vals = {}
        for k, v in r.items():
            if k in ("SKU", "PN"): continue
            try:
                vals[k] = float(v) if v and v.strip() else 0
            except:
                vals[k] = 0
        result[sku] = vals
    return result

def aggregate_weekly(daily, cut_day="Saturday"):
    """Aggregate daily → weekly based on cut_day.
    cut_day: 'Monday'~'Sunday'. 
    Week defined: the 7-day period ending on cut_day.
    Label = the cut_day date of that week.
    """
    if not daily: return {}
    # cut_day → weekday index (Mon=0..Sun=6)
    dow_map = {"Monday": 0, "Tuesday": 1, "Wednesday": 2, "Thursday": 3,
               "Friday": 4, "Saturday": 5, "Sunday": 6}
    target_dow = dow_map.get(cut_day, 5)  # default Saturday

    # Group dates by week (the week whose end day = target_dow)
    weeks = defaultdict(list)
    for d_str in daily:
        try:
            dt = datetime.strptime(d_str, "%Y-%m-%d")
        except: continue
        # Calculate offset to get to the target cut day
        current_dow = dt.weekday()  # Mon=0
        diff = (target_dow - current_dow) % 7
        week_end = dt + timedelta(days=diff)
        weeks[week_end.strftime("%Y-%m-%d")].append(d_str)

    result = {}
    for wk_label, day_list in weeks.items():
        total = sum(daily[d] for d in day_list if daily.get(d))
        if total > 0:
            result[wk_label] = round(total, 0)
    return result

def extract_weekly(daily, cut_day="Saturday"):
    """For cumulative data: take the value on or before the cut_day each week."""
    if not daily: return {}
    dow_map = {"Monday": 0, "Tuesday": 1, "Wednesday": 2, "Thursday": 3,
               "Friday": 4, "Saturday": 5, "Sunday": 6}
    target_dow = dow_map.get(cut_day, 5)

    weeks = defaultdict(list)
    for d_str in daily:
        try:
            dt = datetime.strptime(d_str, "%Y-%m-%d")
        except: continue
        current_dow = dt.weekday()
        diff = (target_dow - current_dow) % 7
        week_end = dt + timedelta(days=diff)
        weeks[week_end.strftime("%Y-%m-%d")].append(d_str)

    result = {}
    for wk_label, day_list in weeks.items():
        # For cumulative, take the latest day's value in the week
        day_list.sort()
        if day_list:
            last = day_list[-1]
            if daily.get(last):
                result[wk_label] = round(daily[last], 0)
    return result

def build_report_data(config=None):
    if config is None:
        config = {"etd_cut": "Saturday", "output_cut": "Wednesday", "exf_cut": "Saturday"}

    sku_attrs, sku_to_gb, sku_pallet, gb_style_color = read_sku_master()

    # Read all wide tables
    plan_gated = read_wide_csv("plan_output_gated.csv")
    plan_ungated = read_wide_csv("plan_output_ungated.csv")
    fcst_wide = read_wide_csv("forecast.csv")
    ctb_wide = read_wide_csv("ctb_cum.csv")

    # Load GB CTB file for GB-level CTB values
    gb_ctb_raw = {}
    gb_ctb_fp = DATA_DIR.parent / "data" / "Modelo GB CTB Publish 0710 final.xlsx"
    if gb_ctb_fp.exists():
        import openpyxl
        wb = openpyxl.load_workbook(gb_ctb_fp, data_only=True)
        ws = wb["Modelo GB CTB "]
        for r in range(4, ws.max_row + 1):
            title = str(ws.cell(r, 1).value or "")
            if title != "CTB": continue
            style = str(ws.cell(r, 6).value or "").strip()
            color = str(ws.cell(r, 7).value or "").strip().upper()
            if not style or not color or color == "COLOR": continue
            daily = {}
            for c in range(10, min(ws.max_column + 1, 420)):
                dt = ws.cell(3, c).value
                v = ws.cell(r, c).value
                if isinstance(dt, datetime) and v is not None and isinstance(v, (int, float)):
                    daily[dt.strftime("%Y-%m-%d")] = v
            # Aggregate across suppliers for same (style, color)
            if (style, color) not in gb_ctb_raw:
                gb_ctb_raw[(style, color)] = defaultdict(float)
            for d, v in daily.items():
                gb_ctb_raw[(style, color)][d] += v
    # Convert GB CTB to weekly
    gb_ctb_weekly = {}
    for (style, color), daily in gb_ctb_raw.items():
        gb_ctb_weekly[(style, color)] = extract_weekly(daily, config["etd_cut"])

    # FG = only SKUs from master mapping (no GB PNs in FG dimension)
    all_skus = sorted(s for s in set(list(plan_gated.keys()) + list(plan_ungated.keys()) +
                                     list(fcst_wide.keys()) + list(ctb_wide.keys()))
                      if s in sku_attrs)

    # Aggregate each plan by configured cut days
    agg = {}
    for label, data, cut in [("GATED_ETD", plan_gated, config["etd_cut"]),
                              ("UNGATED_ETD", plan_ungated, config["etd_cut"]),
                              ("GATED_PACK", plan_gated, config["output_cut"]),
                              ("UNGATED_PACK", plan_ungated, config["output_cut"]),
                              ("FCST", fcst_wide, config["exf_cut"])]:
        agg[label] = {}
        for sku in all_skus:
            daily = data.get(sku, {})
            if label.startswith("FCST"):
                agg[label][sku] = extract_weekly(daily, cut)
            else:
                agg[label][sku] = aggregate_weekly(daily, cut)

    # Collect all week labels
    all_weeks = set()
    for a in agg.values():
        for v in a.values():
            all_weeks.update(v.keys())
    # Also from CTB
    for sku in all_skus:
        if sku in ctb_wide:
            # CTB is cumulative, use etd_cut to extract weekly
            w = extract_weekly(ctb_wide[sku], config["etd_cut"])
            all_weeks.update(w.keys())
    all_weeks = sorted(all_weeks)

    def fill(vals): return {w: vals.get(w, None) for w in all_weeks}
    def diff(b, s):
        ks = set(list(b.keys()) + list(s.keys()))
        return {k: ((b.get(k) or 0) - (s.get(k) or 0)) for k in ks}

    rows = []
    for sku in all_skus:
        attrs = sku_attrs.get(sku, {})
        style = attrs.get("Style", "")
        color = attrs.get("Color", "")
        gb = sku_to_gb.get(sku, "")
        pallet = sku_pallet.get(sku, 864)

        ue = agg["UNGATED_ETD"].get(sku, {})
        up = agg["UNGATED_PACK"].get(sku, {})
        ge = agg["GATED_ETD"].get(sku, {})
        gp = agg["GATED_PACK"].get(sku, {})
        exf = agg["FCST"].get(sku, {})

        # Apply pallet rounding to ETD
        ue = {w: (int(v)//pallet)*pallet for w, v in ue.items() if v and v > 0}
        ge = {w: (int(v)//pallet)*pallet for w, v in ge.items() if v and v > 0}

        # CTB
        ctb = {}
        if sku in ctb_wide:
            ctb = extract_weekly(ctb_wide[sku], config["etd_cut"])

        usage = attrs.get("Usage", "")
        base = {"PN": sku, "Usage": usage, "Style": style, "Color": color,
                "GB_PN": gb, "Pallet_Qty": pallet, "_dim": "FG"}

        def cut(cd): return {"Cut Day": cd}
        rows.append({**base, "Version-Type": "ExF", "Version-Detail": "", **cut(config["exf_cut"]), **fill(exf)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "ETD", **cut(config["etd_cut"]), **fill(ue)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "ETD vs ExF", **cut(config["etd_cut"]), **fill(diff(ue, exf))})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "Packout", **cut(config["output_cut"]), **fill(up)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "Packout vs ExF", **cut(config["output_cut"]), **fill(diff(up, exf))})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "ETD", **cut(config["etd_cut"]), **fill(ge)})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "ETD vs ExF", **cut(config["etd_cut"]), **fill(diff(ge, exf))})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "Packout", **cut(config["output_cut"]), **fill(gp)})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "Packout vs ExF", **cut(config["output_cut"]), **fill(diff(gp, exf))})
        rows.append({**base, "Version-Type": "CTB", "Version-Detail": "", **cut(""), **fill(ctb)})

    # GB dimension
    gb_groups = defaultdict(list)
    for sku in all_skus:
        g = sku_to_gb.get(sku, "")
        if g: gb_groups[g].append(sku)

    sku_data = {(r["PN"], r["Version-Type"], r["Version-Detail"]): {w: r.get(w) for w in all_weeks} for r in rows}

    for gb, skus in gb_groups.items():
        if len(skus) < 2: continue
        sc = gb_style_color.get(gb, ("",""))
        pallet = 864
        base = {"PN": gb, "Usage": "", "Style": sc[0], "Color": sc[1],
                "GB_PN": gb, "Pallet_Qty": pallet, "_dim": "GB"}

        for vt, vd in [("ExF",""), ("Ungated","ETD"), ("Ungated","ETD vs ExF"),
                       ("Ungated","Packout"), ("Ungated","Packout vs ExF"),
                       ("Gated","ETD"), ("Gated","ETD vs ExF"),
                       ("Gated","Packout"), ("Gated","Packout vs ExF"),
                       ("CTB","")]:
            if vt == "CTB":
                # Use GB CTB file data when available
                sc_key = (sc[0], sc[1])
                if sc_key in gb_ctb_weekly:
                    vals = {w: round(gb_ctb_weekly[sc_key].get(w, 0), 0) for w in all_weeks if gb_ctb_weekly[sc_key].get(w)}
                else:
                    vals = {}
                    for w in all_weeks:
                        s = sum(sku_data.get((s, "CTB", ""), {}).get(w, 0) or 0 for s in skus)
                        if s > 0: vals[w] = round(s, 0)
            else:
                vals = {}
                for w in all_weeks:
                    s = sum(sku_data.get((s, vt, vd), {}).get(w, 0) or 0 for s in skus)
                    if s > 0: vals[w] = round(s, 0)
            cd = "" if vt == "CTB" else (config["output_cut"] if "Packout" in vd else config["etd_cut"])
            rows.append({**base, "Version-Type": vt, "Version-Detail": vd, "Cut Day": cd, **fill(vals)})

    return {"rows": rows, "weeks": all_weeks}

if __name__ == "__main__":
    import sys, time
    t0 = time.time()
    config = {"etd_cut": "Saturday", "output_cut": "Wednesday", "exf_cut": "Saturday"}
    data = build_report_data(config)
    t1 = time.time()
    print(f"Built {len(data['rows'])} rows, {len(data['weeks'])} weeks in {t1-t0:.1f}s")
    static_dir = os.path.join(os.path.dirname(__file__), "static")
    os.makedirs(static_dir, exist_ok=True)
    with open(os.path.join(static_dir, "report_data.json"), "w") as f:
        json.dump(data, f, ensure_ascii=False, default=str)
    sz = os.path.getsize(os.path.join(static_dir, "report_data.json"))
    print(f"JSON: {sz/1024/1024:.1f} MB")
    wl = {}
    for w in data['weeks']:
        try:
            dt = datetime.strptime(w, '%Y-%m-%d')
            wk = (dt.day - 1) // 7 + 1
            wl[w] = f"{dt.strftime('%b')} Wk{wk} ({dt.strftime('%b %d')})"
        except: wl[w] = w
    with open(os.path.join(static_dir, "meta.json"), "w") as f:
        json.dump({"week_labels": wl}, f)
    print("Done")
