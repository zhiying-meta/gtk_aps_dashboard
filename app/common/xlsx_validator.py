"""
Unified xlsx validation - previously duplicated in plan_merge, utilization, io
"""
import os
import zipfile


def is_valid_xlsx_by_content(file_path: str) -> bool:
    """Check if file is xlsx by reading magic number PK"""
    try:
        if not os.path.exists(file_path):
            return False
        if os.path.getsize(file_path) < 10:
            return False
        with open(file_path, 'rb') as fh:
            head = fh.read(4)
            return head[:2] == b'PK'
    except Exception:
        return False


def is_valid_xlsx_deep(file_path: str):
    """
    Deep validation: try to open with zipfile + openpyxl to catch truncated files
    Returns (is_valid: bool, message: str)
    """
    try:
        if not os.path.exists(file_path) or os.path.getsize(file_path) < 100:
            return False, "File too small or not exists"
        with open(file_path, 'rb') as fh:
            if fh.read(2) != b'PK':
                return False, "Not a zip/xlsx (no PK header)"
        try:
            with zipfile.ZipFile(file_path, 'r') as zf:
                namelist = zf.namelist()
                if not namelist:
                    return False, "Empty zip"
        except Exception as e:
            return False, f"Bad zip: {e}"
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            _ = wb.sheetnames
            wb.close()
        except EOFError as e:
            return False, f"Truncated/corrupted (EOFError): {e} — file may have been incompletely uploaded"
        except Exception as e:
            err = str(e).lower()
            if "eof" in err or "truncated" in err or "not a zip" in err or "badzip" in err:
                return False, f"Corrupted xlsx ({e})"
        return True, "OK"
    except Exception as e:
        return False, f"Validation failed: {e}"


def validate_xlsx_openpyxl(file_path: str):
    """Validate via openpyxl open - returns (ok, error_msg)"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        wb.close()
        return True, ""
    except Exception as e:
        return False, str(e)
