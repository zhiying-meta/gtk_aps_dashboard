"""
Unified template/xlsx buffer creation - previously duplicated
"""
import io
import zipfile


def xlsx_buf(header, rows=None, sheet_name="Sheet1"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(header)
    for r in rows or []:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_buf_multi_sheet(sheets: dict):
    """
    sheets: {sheet_name: (header, rows)}
    Returns BytesIO
    """
    import openpyxl
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, (header, rows) in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        ws.append(header)
        for r in rows or []:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def zip_files(file_map: dict):
    """
    file_map: {arcname: file_path or BytesIO or bytes}
    Returns BytesIO of zip
    """
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w", zipfile.ZIP_DEFLATED) as zf:
        for arcname, content in file_map.items():
            if isinstance(content, str):  # file path
                zf.write(content, arcname=arcname)
            elif isinstance(content, io.BytesIO):
                zf.writestr(arcname, content.getvalue())
            elif isinstance(content, bytes):
                zf.writestr(arcname, content)
            else:
                # assume file-like
                zf.writestr(arcname, content)
    zb.seek(0)
    return zb
