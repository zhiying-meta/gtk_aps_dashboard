"""
Server: static files + API endpoints
"""
import os, json, sys, io
from flask import Flask, request, jsonify, send_from_directory, send_file
from datetime import datetime

sys.path.insert(0, os.path.dirname(__file__))
from backend import build_report_data
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

app = Flask(__name__, static_folder="static")

LABEL_MAP = {0:"Mon",1:"Tue",2:"Wed",3:"Thu",4:"Fri",5:"Sat",6:"Sun"}

def fmt_week_label(w):
    try:
        dt = datetime.strptime(w, "%Y-%m-%d")
        wk = (dt.day - 1) // 7 + 1
        return f"{dt.strftime('%b')} Wk{wk} ({dt.strftime('%b %d')})"
    except:
        return w

@app.after_request
def no_cache(resp):
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return resp

@app.route("/")
def index():
    return send_from_directory("static", "index.html")

@app.route("/<path:path>")
def static_files(path):
    return send_from_directory("static", path)

@app.route("/api/data", methods=["POST"])
def api_data():
    config = request.get_json() or {}
    cfg = {
        "etd_cut": config.get("etd_cut", "Saturday"),
        "output_cut": config.get("output_cut", "Wednesday"),
        "exf_cut": config.get("exf_cut", "Saturday"),
    }
    try:
        data = build_report_data(cfg)
        wl = {w: fmt_week_label(w) for w in data["weeks"]}
        return jsonify({"rows": data["rows"], "weeks": data["weeks"], "week_labels": wl, "config": cfg})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/download", methods=["POST"])
def api_download():
    config = request.get_json() or {}
    cfg = {
        "etd_cut": config.get("etd_cut", "Saturday"),
        "output_cut": config.get("output_cut", "Wednesday"),
        "exf_cut": config.get("exf_cut", "Saturday"),
    }
    try:
        data = build_report_data(cfg)
        rows, weeks = data["rows"], data["weeks"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"

        hf = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        hfl = PatternFill("solid", fgColor="1E293B")
        hb = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
        cf = Font(name="微软雅黑", size=10)
        cb = Border(left=Side(style='thin'), right=Side(style='thin'),
                     top=Side(style='thin'), bottom=Side(style='thin'))
        nf = '#,##0'
        fills = {"ExF": PatternFill("solid", fgColor="F0F9FF"),
                 "Ungated": PatternFill("solid", fgColor="F0FDF4"),
                 "Gated": PatternFill("solid", fgColor="FFFBEB"),
                 "CTB": PatternFill("solid", fgColor="FAF5FF")}

        # Fixed columns (same order as ALL_COLS in frontend)
        fixed_keys = ["_dim", "PN", "Usage", "Style", "Color", "Version-Type", "Version-Detail", "Cut Day", "Pallet_Qty"]
        fixed_labels = ["Dim", "PN", "Usage", "Style", "Color", "Version-Type", "Version-Detail", "Cut Day", "Pallet"]

        # Header
        for ci, label in enumerate(fixed_labels + [fmt_week_label(w) for w in weeks], 1):
            cell = ws.cell(1, ci, label)
            cell.font = hf; cell.fill = hfl; cell.border = hb
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 28
        # auto-filter skipped to avoid format issues

        # Data
        last_pn = None
        for ri, r in enumerate(rows, 2):
            vals = []
            for k in fixed_keys:
                if k == "_dim": vals.append(r.get("_dim", ""))
                elif k == "Pallet_Qty": vals.append(r.get("Pallet_Qty") or "")
                else: vals.append(str(r.get(k, "") or ""))
            for w in weeks:
                v = r.get(w)
                vals.append(v if v is not None else "")

            for ci, val in enumerate(vals, 1):
                cell = ws.cell(ri, ci, val)
                cell.font = cf; cell.border = cb
                cell.alignment = Alignment(vertical='center')
                if ci > len(fixed_keys) and isinstance(val, (int, float)):
                    cell.number_format = nf
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    if val > 0: cell.font = Font(name="微软雅黑", size=10, color="059669")
                    elif val < 0: cell.font = Font(name="微软雅黑", size=10, color="DC2626")

            # Type background
            ft = fills.get(r.get("Version-Type"))
            if ft:
                for ci in range(1, len(fixed_labels) + len(weeks) + 1):
                    ws.cell(ri, ci).fill = ft

            # PN border
            is_new = r.get("PN") != last_pn
            if is_new and ri > 2:
                for ci in range(1, len(fixed_labels) + len(weeks) + 1):
                    c = ws.cell(ri, ci)
                    c.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                      top=Side(style='medium', color="94A3B8"),
                                      bottom=Side(style='thin'))
            last_pn = r.get("PN")

        ws.freeze_panes = None  # skip freeze to avoid format issues

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        from flask import Response
        return Response(
            buf.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=report.xlsx"}
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8501))
    print(f"🌐 Server at http://localhost:{port}")
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
