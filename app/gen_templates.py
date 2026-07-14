"""
Generate single xlsx template with 6 sheets + schema JSON for frontend
"""
import openpyxl, os, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")
os.makedirs(TEMPLATE_DIR, exist_ok=True)

HDR_Fill = PatternFill("solid", fgColor="1E293B")
HDR_Font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
C_Font = Font(name="微软雅黑", size=10)
THIN = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))

SCHEMA = {}

def add_sheet(wb, sheet_name, headers, schema_fields, schema_note):
    ws = wb.create_sheet(title=sheet_name)
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = HDR_Font; c.fill = HDR_Fill; c.border = THIN
        c.alignment = Alignment(horizontal='center', vertical='center')
    SCHEMA[sheet_name] = {"fields": schema_fields, "note": schema_note}
    print(f"  Sheet: {sheet_name} ({len(headers)} cols)")

# Create template workbook
wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

add_sheet(wb, "sku_master",
    headers=["SKU","Usage","Style","Color","GB_PN","FR_PN","LT_PN","RT_PN","Pallet_Qty"],
    schema_fields=[
        ("SKU","String","物料编码","SK-1001879-01"),
        ("Usage","String","用途：MP/Dummy/Demo","MP"),
        ("Style","String","眼镜款式","Rectangle M"),
        ("Color","String","颜色（大写）","BLACK"),
        ("GB_PN","String","GB维度物料编码","GB-Rec M-BLACK"),
        ("FR_PN","String","Frame物料编码","FR-Rec M-BLACK"),
        ("LT_PN","String","Left Temple物料编码","LT-Rec M-BLACK"),
        ("RT_PN","String","Right Temple物料编码","RT-Rec M-BLACK"),
        ("Pallet_Qty","Integer","托盘数量（默认864）","864"),
    ],
    schema_note="SKU主数据表，必填。每个SKU一行，Pallet_Qty不填则默认864。")

add_sheet(wb, "plan_output_gated",
    headers=["PN","2026-04-16","2026-04-17","2026-04-18","..."],
    schema_fields=[
        ("PN","String","物料编码（SKU或GB/FR/LT/RT PN）","SK-1001879-01"),
        ("2026-04-16","Integer","日级OUTPUT产出量","0"),
        ("...","...","后续列为每天一列，格式yyyy-MM-dd","..."),
    ],
    schema_note="Gated版本日级排产产出。第一列PN，后续列为日期(yyyy-MM-dd)，值为当日产出量。日期列数不限。")

add_sheet(wb, "plan_output_ungated",
    headers=["PN","2026-04-16","2026-04-17","2026-04-18","..."],
    schema_fields=[
        ("PN","String","物料编码（SKU或GB/FR/LT/RT PN）","SK-1001879-01"),
        ("2026-04-16","Integer","日级OUTPUT产出量","0"),
        ("...","...","后续列为每天一列","..."),
    ],
    schema_note="Ungated版本日级排产产出。格式同plan_output_gated。")

add_sheet(wb, "forecast",
    headers=["SKU","2026-04-04","2026-04-11","2026-04-18","..."],
    schema_fields=[
        ("SKU","String","物料编码（仅SKU级）","SK-1001879-01"),
        ("2026-04-04","Integer","Forecast周值（周六截止）","0"),
        ("...","...","后续列为每周六日期","..."),
    ],
    schema_note="周级Forecast数据。列头为周六日期，值为该周Forecast量。")

add_sheet(wb, "ctb_sku_cum",
    headers=["SKU","2026-02-04","2026-02-05","2026-02-06","..."],
    schema_fields=[
        ("SKU","String","物料编码（SKU级）","SK-1001879-01"),
        ("2026-02-04","Integer","日级CTB累计值","0"),
        ("...","...","后续列为每天一列","..."),
    ],
    schema_note="SKU级CTB累计值。日级数据，值为截止该日的CTB累计量（非日增量）。")

add_sheet(wb, "ctb_gb_cum",
    headers=["PN","2026-02-04","2026-02-05","2026-02-06","..."],
    schema_fields=[
        ("PN","String","GB物料编码","GB-Rec M-BLACK"),
        ("2026-02-04","Integer","日级GB CTB累计值","0"),
        ("...","...","后续列为每天一列","..."),
    ],
    schema_note="GB级CTB累计值。可选，如无GB级CTB数据可不填。")

template_path = os.path.join(TEMPLATE_DIR, "input_template.xlsx")
wb.save(template_path)
print(f"\nSaved: {template_path}")

# Schema JSON
with open(os.path.join(TEMPLATE_DIR, "schema.json"), "w", encoding="utf-8") as f:
    json.dump(SCHEMA, f, ensure_ascii=False, indent=2)
print("Schema JSON saved")
