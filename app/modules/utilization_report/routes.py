import os
import shutil
import tempfile
import pathlib

from flask import Blueprint, jsonify, request

from app.modules.utilization_report.engine import (
    get_all_caches,
    get_cache_for_version,
    get_status,
    build_report,
    build_compare_report,
    reload_all,
    _compute_records,
)

util_bp = Blueprint("utilization_report", __name__)

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
DEFAULT_DATA_DIR = os.path.join(PROJECT_ROOT, "data")

# For upload, we expect 2 files: calendar and schedule
# TARGET mapping similar to io_report but for utilization

def _json_error(msg, code=500):
    return jsonify({"error": msg}), code

@util_bp.route("/api/utilization/status", methods=["GET"])
def api_status():
    try:
        from app.modules.utilization_report.engine import _CACHES
        import pathlib
        base_path = pathlib.Path(DEFAULT_DATA_DIR)

        # First, validate cache against filesystem — if user manually deletes gated/ungated folders,
        # cache should be invalidated and not returned as Ready
        # This explains: manually clear gated and ungated, but matrix still shows other version with no actual data
        for k in list(_CACHES.keys()):
            if k.startswith(DEFAULT_DATA_DIR + "::"):
                ver = k.split("::")[-1]
                util_dir = base_path / "utilization" / ver
                cal = util_dir / "工作日历快照.xlsx"
                sched = util_dir / "排产结果表.xlsx"
                if not (util_dir.exists() and cal.exists() and sched.exists()):
                    _CACHES.pop(k, None)

        cached_versions = {}
        for k in _CACHES.keys():
            if k.startswith(DEFAULT_DATA_DIR + "::"):
                ver = k.split("::")[-1]
                c = _CACHES[k]
                cached_versions[ver] = {
                    'lines': len(c.lines),
                    'dates': len(c.dates),
                    'shifts': len(c.shifts),
                    'records_shift': len(c.records_shift),
                    'records_day': len(c.records_day),
                }

        files_found = []
        file_details = {}  # ver -> list of {name, size_kb, exists, is_calendar, is_schedule}
        file_summary = {}  # ver -> summary string like "calendar:xxx + schedule:yyy"
        for ver in ['gated','ungated']:
            util_dir = base_path / "utilization" / ver
            cal_path = util_dir / "工作日历快照.xlsx"
            sched_path = util_dir / "排产结果表.xlsx"
            ver_files = []
            if cal_path.exists():
                try:
                    sz = cal_path.stat().st_size
                    ver_files.append({"name": "工作日历快照.xlsx", "original": "工作日历快照.xlsx", "type": "calendar", "size": sz, "size_kb": round(sz/1024,1), "exists": True})
                except:
                    ver_files.append({"name": "工作日历快照.xlsx", "type": "calendar", "exists": True, "size_kb": 0})
            if sched_path.exists():
                try:
                    sz = sched_path.stat().st_size
                    ver_files.append({"name": "排产结果表.xlsx", "original": "排产结果表.xlsx", "type": "schedule", "size": sz, "size_kb": round(sz/1024,1), "exists": True})
                except:
                    ver_files.append({"name": "排产结果表.xlsx", "type": "schedule", "exists": True, "size_kb": 0})
            # Also list any other xlsx/zip files in dir for debugging (like original named uploads if user copied manually) – IO-style, only xlsx/zip, ignore pkl/cache
            if util_dir.exists():
                try:
                    for f in util_dir.iterdir():
                        if f.is_file() and f.name not in ("工作日历快照.xlsx", "排产结果表.xlsx"):
                            # Show extra files as unknown only if xlsx or zip (like IO shows only relevant files)
                            if f.suffix.lower() not in (".xlsx", ".zip"):
                                continue
                            # Skip cache files
                            if f.name.lower().endswith('.pkl') or f.name.startswith('cache_'):
                                continue
                            try:
                                sz = f.stat().st_size
                            except:
                                sz = 0
                            ver_files.append({"name": f.name, "original": f.name, "type": "unknown", "size": sz, "size_kb": round(sz/1024,1), "exists": True})
                except:
                    pass
            if ver_files:
                file_details[ver] = ver_files
                # Build summary
                cals = [x for x in ver_files if x.get('type')=='calendar']
                scheds = [x for x in ver_files if x.get('type')=='schedule']
                if cals and scheds:
                    file_summary[ver] = f"calendar={cals[0]['name']}({cals[0].get('size_kb',0)}KB) + schedule={scheds[0]['name']}({scheds[0].get('size_kb',0)}KB)"
                elif cals:
                    file_summary[ver] = f"calendar={cals[0]['name']}({cals[0].get('size_kb',0)}KB) only — missing schedule"
                elif scheds:
                    file_summary[ver] = f"schedule={scheds[0]['name']}({scheds[0].get('size_kb',0)}KB) only — missing calendar"
                else:
                    file_summary[ver] = ", ".join([f['name'] for f in ver_files])

            if util_dir.exists() and cal_path.exists() and sched_path.exists():
                files_found.append(ver)

        # No fallback to IVY folders — after clear all, files_found should be empty, so UI shows Not Ready and matrix empty
        # Only Ready if files exist AND cache exists (validated above)

        # Truly ready only when cache exists and files exist
        if cached_versions:
            return jsonify({
                "loaded": True,
                "versions": list(cached_versions.keys()),
                "details": cached_versions,
                "files_found": files_found,
                "file_details": file_details,
                "file_summary": file_summary,
                "detected_message": f"✅ Utilization: detected {sum(len(v) for v in file_details.values())} files — " + "; ".join([f"{ver}: {file_summary.get(ver,'')}" for ver in files_found]) if files_found else "No files",
            })
        else:
            # Not yet computed, but files may exist
            return jsonify({
                "loaded": False,
                "versions": [],
                "details": {},
                "files_found": files_found,
                "file_details": file_details,
                "file_summary": file_summary,
                "needs_compute": len(files_found) > 0,
                "detected_message": f"⚠️ Files found but not computed: " + "; ".join([f"{ver}: {file_summary.get(ver,'')}" for ver in file_details.keys()]) if file_details else "No files",
            })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(str(e))

@util_bp.route("/api/utilization/meta", methods=["GET"])
def api_meta():
    try:
        caches = get_all_caches(DEFAULT_DATA_DIR)
        if not caches:
            return _json_error("No utilization data loaded. Upload calendar + schedule or place files in data/ folder.", 404)
        # Use gated as primary for meta, fallback to any
        primary = caches.get('gated') or list(caches.values())[0]
        return jsonify({
            "lines": primary.lines,
            "dates": primary.dates,
            "shifts": primary.shifts,
            "versions": list(caches.keys()),
            "date_min": min(primary.dates) if primary.dates else "",
            "date_max": max(primary.dates) if primary.dates else "",
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(str(e))

@util_bp.route("/api/utilization/reports", methods=["GET"])
def api_reports():
    try:
        caches = get_all_caches(DEFAULT_DATA_DIR)
        if not caches:
            return _json_error("No data. Upload first.", 404)

        mode = request.args.get("mode", "shift")  # shift or day
        version = request.args.get("version", "gated")  # gated / ungated / compare
        line_filter = request.args.get("line_code", "") or request.args.get("line", "")
        shift_filter = request.args.get("shift_name", "") or request.args.get("shift", "")
        date_from = request.args.get("date_from", "")
        date_to = request.args.get("date_to", "")

        filters = {
            "line_filter": line_filter,
            "shift_filter": shift_filter,
            "date_from": date_from,
            "date_to": date_to,
        }

        if version == "compare":
            if 'gated' in caches and 'ungated' in caches:
                gated = caches['gated']
                ungated = caches['ungated']
                data = build_compare_report(gated, ungated, mode=mode, **filters)
                return jsonify({
                    "mode": mode,
                    "version": "compare",
                    "count": len(data),
                    "records": data[:2000],  # limit
                    "total": len(data),
                })
            else:
                return _json_error("Need both gated and ungated data for compare. Currently have: " + ",".join(caches.keys()), 400)
        else:
            # Single version
            cache = caches.get(version)
            if not cache:
                # fallback to first available
                cache = list(caches.values())[0]
                version = cache.version
            data = build_report(cache, mode=mode, **filters)
            return jsonify({
                "mode": mode,
                "version": cache.version,
                "count": len(data),
                "records": data[:2000],
                "total": len(data),
                "lines": cache.lines[:100],
                "shifts": cache.shifts,
            })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(str(e))

@util_bp.route("/api/utilization/pivot", methods=["GET"])
def api_pivot():
    """
    Pivot table for utilization matrix like Packout report.
    Supports version_type per row: Gated / Ungated
    mode=shift -> columns = date + shift
    mode=day -> columns = date
    Returns rows with line_code + version_type, plus columns for dates.
    Includes color coding, thick line separation by line_code.
    """
    try:
        caches = get_all_caches(DEFAULT_DATA_DIR)
        if not caches:
            # After clear all, no caches — return empty with not_ready flag so matrix shows no values (as user expects)
            return jsonify({
                "mode": request.args.get("mode", "day"),
                "versions": [],
                "requested_version": request.args.get("version", "all"),
                "not_ready": True,
                "columns": [],
                "rows": [],
                "detail": {},
                "lines": [],
                "total_lines": 0,
                "total_cols": 0,
                "message": "No data — both Gated and Ungated are Not Ready (empty). Upload at least one module.",
                "all_cleared": True
            })

        mode = request.args.get("mode", "day")
        version_param = request.args.get("version", "all")  # gated, ungated, all, compare
        version_filter = request.args.get("version_type", "")  # optional filter Gated/Ungated
        line_filter = request.args.get("line_code", "") or request.args.get("line", "")
        date_from = request.args.get("date_from", "")
        date_to = request.args.get("date_to", "")

        from app.modules.utilization_report.engine import build_report

        # Determine which versions to include
        versions_to_include = []
        if version_param == "all" or version_param == "compare" or "," in version_param:
            # Include all available caches - if only gated ready, show gated only (ungated empty)
            if version_param == "all" or version_param == "compare":
                versions_to_include = list(caches.keys())
            else:
                # e.g. version="gated,ungated" but filter only those that exist -> shows ready ones, missing ones remain not ready
                versions_to_include = [v.strip() for v in version_param.split(",") if v.strip() in caches]
                # If none of requested versions exist, return not_ready empty (explicit Not Ready display)
                if not versions_to_include:
                    return jsonify({
                        "mode": mode,
                        "versions": [],
                        "requested_version": version_param,
                        "not_ready": True,
                        "columns": [],
                        "rows": [],
                        "detail": {},
                        "lines": [],
                        "total_lines": 0,
                        "total_cols": 0,
                        "message": f"{version_param} not ready — no data uploaded, other version may be ready"
                    })
        else:
            # Single version explicit
            if version_param in caches:
                versions_to_include = [version_param]
            else:
                # Version not ready -> return empty with not_ready flag so frontend can show "Not Ready" explicitly
                return jsonify({
                    "mode": mode,
                    "versions": [],
                    "requested_version": version_param,
                    "not_ready": True,
                    "columns": [],
                    "rows": [],
                    "detail": {},
                    "lines": [],
                    "total_lines": 0,
                    "total_cols": 0,
                    "message": f"{version_param} not ready — no data uploaded"
                })

        # If version_type filter provided, further filter
        if version_filter:
            vt_low = version_filter.lower()
            # Map Gated/Ungated to cache keys
            filtered = []
            for ver in versions_to_include:
                if vt_low in ver.lower() or vt_low in ver:
                    filtered.append(ver)
            if filtered:
                versions_to_include = filtered

        # Collect records per version
        all_recs = []  # list of dict with extra version_type
        for ver in versions_to_include:
            cache = caches.get(ver)
            if not cache:
                continue
            recs = build_report(cache, mode=mode, line_filter=line_filter, date_from=date_from, date_to=date_to)
            for r in recs:
                nr = dict(r)
                nr['version_type'] = ver.capitalize()  # Gated / Ungated
                nr['version_key'] = ver
                all_recs.append(nr)

        if not all_recs:
            return jsonify({"mode": mode, "columns": [], "rows": [], "detail": {}, "lines": [], "total_lines": 0, "total_cols": 0})

        # Build columns
        if mode == "shift":
            cols_set = set()
            for r in all_recs:
                col = f"{r['plan_date']}|{r['shift_name']}"
                cols_set.add(col)
            cols = sorted(list(cols_set))
        else:
            cols = sorted(list(set(r['plan_date'] for r in all_recs)))

        # Limit columns for performance but enough to cover ~1 year
        # Day mode: ~400 days, Shift mode: ~800 cols (400 days*2), set to 1000 to cover full range
        max_cols = 1000
        total_cols_before = len(cols)
        truncated = False
        if len(cols) > max_cols:
            cols = cols[:max_cols]
            truncated = True

        # Build matrix keyed by (line_code, version_type) -> col -> util
        # Also keep detail
        from collections import defaultdict
        matrix = {}  # key = (line, version_type) -> dict col->util
        detail = {}  # key = (line, version_type) -> col -> detail
        lines_set = set()
        for r in all_recs:
            key = (r['line_code'], r['version_type'])
            lines_set.add(r['line_code'])
            if key not in matrix:
                matrix[key] = {}
                detail[key] = {}
            col = f"{r['plan_date']}|{r['shift_name']}" if mode == "shift" else r['plan_date']
            if col not in cols:
                continue
            util_pct = r.get('utilization_pct', 0)
            capped = min(100.0, util_pct) if util_pct is not None else 0
            matrix[key][col] = capped
            detail[key][col] = {
                'util_raw': util_pct,
                'util_capped': capped,
                'load': r['load'],
                'capacity': r['capacity'],
                'uph': r['uph'],
                'efficiency': r['efficiency'],
                'working_hours': r['working_hours'],
                'is_overload': r.get('is_overload', False),
            }

        # Sort keys by line_code, then version_type (Gated before Ungated for consistency)
        def sort_key(k):
            line, vtype = k
            # Gated first
            order = 0 if vtype.lower() == 'gated' else 1
            return (line, order, vtype)
        sorted_keys = sorted(matrix.keys(), key=sort_key)

        # Build rows for frontend with version_type field and thick border flag
        rows = []
        last_line = None
        for (line, vtype) in sorted_keys:
            row = {
                'line_code': line,
                'version_type': vtype,
                'is_new_line': line != last_line,
            }
            last_line = line
            for col in cols:
                row[col] = matrix[(line, vtype)].get(col, None)
            rows.append(row)

        # Convert detail to string keys for JSON (line|version_type)
        detail_json = {}
        for (line, vtype), col_map in detail.items():
            key_str = f"{line}||{vtype}"
            detail_json[key_str] = col_map

        return jsonify({
            "mode": mode,
            "versions": versions_to_include,
            "columns": cols,
            "rows": rows,
            "detail": detail_json,
            "lines": sorted(list(lines_set)),
            "total_lines": len(sorted_keys),
            "total_cols": len(cols),
            "total_cols_before": total_cols_before,
            "truncated": truncated,
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(str(e))

@util_bp.route("/api/utilization/upload", methods=["POST"])
def api_upload():
    """
    Upload handler – IO-style (reference IO report import):
    - Accepts: multiple xlsx, zip, folder (webkitdirectory) for gated/ungated
    - Auto-classify by filename keywords (calendar=工作日历, schedule=排产结果) + version (gated/ungated)
    - Informs specifically which files were uploaded (like IO: detected N files — ...)
    - Supports:
      * per-version one-click: ?version=gated|ungated (calendar+schedule or zip)
      * separate upload: calendar alone or schedule alone
      * one-click: calendar+schedule together for single version
      * multi-version: gated+ungated together (4 files or zip with both)
    Returns detailed file list for frontend persistent badge.
    """
    if not request.files:
        return _json_error("Expected calendar and schedule files (xlsx or zip)", 400)

    forced_version = (request.args.get('version') or request.form.get('version') or '').strip().lower()
    if forced_version not in ('gated', 'ungated'):
        forced_version = None

    tmp_dir = tempfile.mkdtemp(prefix="util_upload_")
    tmp_extract_dir = tempfile.mkdtemp(prefix="util_extract_")
    try:
        import zipfile

        # ---------- IO-style collection ----------
        all_files = []
        for key in request.files:
            all_files.extend([(key, f) for f in request.files.getlist(key)])

        if not all_files:
            return _json_error("No files uploaded", 400)

        # Helper: check xlsx by content PK header
        def _is_valid_xlsx_content(p):
            try:
                if not os.path.exists(p) or os.path.getsize(p) < 10:
                    return False
                with open(p, 'rb') as fh:
                    return fh.read(2) == b'PK'
            except:
                return False

        def _is_valid_xlsx_deep(p):
            """Deep validation: try to open with zipfile + openpyxl to catch truncated files (EOFError)"""
            try:
                if not os.path.exists(p) or os.path.getsize(p) < 100:
                    return False, "File too small or not exists"
                # Check PK header first
                with open(p, 'rb') as fh:
                    if fh.read(2) != b'PK':
                        return False, "Not a zip/xlsx (no PK header)"
                # Try zipfile
                import zipfile
                try:
                    with zipfile.ZipFile(p, 'r') as zf:
                        # Try to list and read central directory
                        namelist = zf.namelist()
                        if not namelist:
                            return False, "Empty zip"
                except Exception as e:
                    return False, f"Bad zip: {e}"
                # Try openpyxl read_only to catch EOFError as in user report
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                    # Try to read at least sheet names
                    _ = wb.sheetnames
                    wb.close()
                except EOFError as e:
                    return False, f"Truncated/corrupted (EOFError): {e} — file may have been incompletely uploaded"
                except Exception as e:
                    err = str(e).lower()
                    if "eof" in err or "truncated" in err or "not a zip" in err or "badzip" in err:
                        return False, f"Corrupted xlsx ({e})"
                    # For other openpyxl errors, still consider valid if zip ok (let pandas handle later)
                return True, "OK"
            except Exception as e:
                return False, f"Validation failed: {e}"

        def _extract_zip_to_tmp(zip_path, out_dir):
            extracted = []
            try:
                with zipfile.ZipFile(zip_path, 'r') as zf:
                    for info in zf.infolist():
                        if info.is_dir():
                            continue
                        # Be permissive: extract xlsx or files containing keywords
                        base = os.path.basename(info.filename)
                        if not base or base.startswith('.') or base.startswith('~$'):
                            continue
                        low = base.lower()
                        # Only care about xlsx or files with our keywords (calendar/schedule)
                        if not (low.endswith('.xlsx') or any(kw in low for kw in ["工作日历", "日历", "calendar", "排产", "schedule"])):
                            # Also extract if size >1KB and not obviously non-excel
                            if info.file_size < 1024 or low.endswith(('.txt','.csv','.pdf')):
                                continue
                        target = os.path.join(out_dir, base)
                        c = 1
                        b, e = os.path.splitext(target)
                        while os.path.exists(target):
                            target = f"{b}_{c}{e}"
                            c += 1
                        with zf.open(info) as src, open(target, 'wb') as dst:
                            shutil.copyfileobj(src, dst)
                        if _is_valid_xlsx_content(target):
                            extracted.append(target)
                        else:
                            # Keep if name matches expected
                            if any(kw in low for kw in ["工作日历","日历","calendar","排产","schedule"]):
                                extracted.append(target)
            except Exception as e:
                print(f"[util] zip extract failed {e}")
            return extracted

        pending_regular = []
        for field_key, f in all_files:
            if not f.filename:
                continue
            fname = f.filename
            if fname.lower().endswith('.zip'):
                zip_tmp = os.path.join(tmp_dir, f"_upload_{field_key}_{os.path.basename(fname)}")
                f.save(zip_tmp)
                extracted = _extract_zip_to_tmp(zip_tmp, tmp_extract_dir)
                for ep in extracted:
                    pending_regular.append((field_key, ep, True))
                try:
                    os.remove(zip_tmp)
                except:
                    pass
            else:
                safe_name = os.path.basename(fname)
                if not safe_name:
                    continue
                # Strip spaces and keep
                safe_name = safe_name.strip()
                temp_path = os.path.join(tmp_extract_dir, f"_raw_{field_key}_{safe_name}")
                c = 1
                b, e = os.path.splitext(temp_path)
                while os.path.exists(temp_path):
                    temp_path = f"{b}_{c}{e}"
                    c += 1
                f.save(temp_path)
                pending_regular.append((field_key, temp_path, True))

        # Second pass: classify and copy to target with IO-style naming
        # For utilization, target filenames are fixed per version: 工作日历快照.xlsx and 排产结果表.xlsx inside version folder
        # But for detection, we keep files in tmp_extract_dir
        final_files = []
        # We will also keep a list of uploaded file info for response
        uploaded_file_info = []  # list of dict with original name, size, detected type

        for field_key, src_path, is_path in pending_regular:
            if not os.path.exists(src_path):
                continue
            base = os.path.basename(src_path)
            # Remove _raw_ prefix for display – format is _raw_{field_key}_{safe_name}
            display_name = base
            if base.startswith("_raw_"):
                prefix = f"_raw_{field_key}_"
                if base.startswith(prefix):
                    display_name = base[len(prefix):]
                else:
                    # Fallback: split _raw_ + field + original
                    # _raw_{field_key}_{safe_name} -> safe_name is after second underscore
                    try:
                        # Remove leading _raw_
                        remainder = base[5:]  # after _raw_
                        # remainder = {field_key}_{safe_name}
                        # Find first underscore after field_key
                        idx = remainder.find("_")
                        if idx != -1:
                            display_name = remainder[idx+1:]
                        else:
                            display_name = remainder
                    except:
                        display_name = base[5:]

            sz = os.path.getsize(src_path)
            # IO-style permissive: keep all files with size >0, even if not PK, to allow proper error reporting later
            # Only skip if size 0 or tiny and name doesn't contain relevant keywords
            if sz < 10:
                continue
            if not _is_valid_xlsx_content(src_path):
                # If not valid xlsx, still keep if name contains relevant keywords or size >1KB (could be old xls or corrupted, let later validation report)
                low = display_name.lower()
                # Keep if contains keywords or size >1KB (permissive)
                if not (any(kw in low for kw in ["工作日历","日历","calendar","排产","schedule","gated","ungated","xlsx","xls"]) or sz > 1024):
                    continue

            final_files.append(src_path)
            uploaded_file_info.append({
                "field": field_key,
                "original": display_name,
                "size": sz,
                "size_kb": round(sz/1024, 1),
                "path": src_path,
            })

        if not final_files:
            existing = os.listdir(tmp_extract_dir)
            return _json_error(f"No xlsx files found (including inside zip). Got {existing}. Expected 工作日历快照.xlsx and 排产结果表.xlsx for gated/ungated.", 400)

        # Classification helpers – IO-style (simplified, no pandas for speed, like IO's _classify_upload)
        def classify_content_type(path):
            name = os.path.basename(path).lower()
            # Also check original display name if path is _raw_
            if name.startswith("_raw_"):
                # Format _raw_{field}_{safe_name} -> extract safe_name
                # Find prefix _raw_{field}_
                try:
                    remainder = name[5:]  # after _raw_
                    idx = remainder.find("_")
                    if idx != -1:
                        name = remainder[idx+1:].lower()
                    else:
                        name = remainder.lower()
                except:
                    name = name[5:].lower()
            is_calendar = False
            is_schedule = False
            if "工作日历" in name or "calendar" in name or "日历" in name:
                is_calendar = True
            elif "排产结果" in name or "schedule" in name or "排产" in name:
                is_schedule = True
            else:
                # Fallback: if name contains uph or capacity keywords? Assume calendar if unknown? But for upload robustness, check file size? No, default to unknown and let caller assign to empty slot
                # To avoid pandas overhead (which can be slow and fail for large files), we use simple heuristic:
                # If file larger than typical schedule? No, use first file as calendar, second as schedule in caller
                is_calendar = False
                is_schedule = False
            return is_calendar, is_schedule

        def classify_file(path):
            name = os.path.basename(path).lower()
            if name.startswith("_raw_"):
                name = "_".join(name.split("_")[2:]).lower() if "_" in name else name[5:].lower()
            is_cal, is_sched = classify_content_type(path)
            ver = None
            if "ungated" in name or "ungate" in name:
                ver = "ungated"
            elif "gated" in name:
                ver = "gated"
            return ver, is_cal, is_sched

        # Group
        if forced_version:
            version_groups = {
                forced_version: {"calendar": None, "schedule": None, "files": []},
            }
            for p in final_files:
                is_cal, is_sched = classify_content_type(p)
                if is_cal and not version_groups[forced_version]["calendar"]:
                    version_groups[forced_version]["calendar"] = p
                    version_groups[forced_version]["files"].append(p)
                elif is_sched and not version_groups[forced_version]["schedule"]:
                    version_groups[forced_version]["schedule"] = p
                    version_groups[forced_version]["files"].append(p)
                else:
                    if not version_groups[forced_version]["calendar"]:
                        version_groups[forced_version]["calendar"] = p
                        version_groups[forced_version]["files"].append(p)
                    elif not version_groups[forced_version]["schedule"]:
                        version_groups[forced_version]["schedule"] = p
                        version_groups[forced_version]["files"].append(p)
        else:
            version_groups = {
                "gated": {"calendar": None, "schedule": None, "files": []},
                "ungated": {"calendar": None, "schedule": None, "files": []},
            }
            unassigned = []
            for p in final_files:
                ver, is_cal, is_sched = classify_file(p)
                if ver and is_cal and not version_groups[ver]["calendar"]:
                    version_groups[ver]["calendar"] = p
                    version_groups[ver]["files"].append(p)
                elif ver and is_sched and not version_groups[ver]["schedule"]:
                    version_groups[ver]["schedule"] = p
                    version_groups[ver]["files"].append(p)
                else:
                    unassigned.append((p, is_cal, is_sched, ver))

            for p, is_cal, is_sched, ver_hint in unassigned:
                target_vers = ["gated", "ungated"] if ver_hint is None else [ver_hint]
                assigned = False
                for tv in target_vers:
                    if is_cal and not version_groups[tv]["calendar"]:
                        version_groups[tv]["calendar"] = p
                        version_groups[tv]["files"].append(p)
                        assigned = True
                        break
                    if is_sched and not version_groups[tv]["schedule"]:
                        version_groups[tv]["schedule"] = p
                        version_groups[tv]["files"].append(p)
                        assigned = True
                        break
                if not assigned:
                    for tv in ["gated", "ungated"]:
                        if is_cal and not version_groups[tv]["calendar"]:
                            version_groups[tv]["calendar"] = p
                            version_groups[tv]["files"].append(p)
                            assigned = True
                            break
                        if is_sched and not version_groups[tv]["schedule"]:
                            version_groups[tv]["schedule"] = p
                            version_groups[tv]["files"].append(p)
                            assigned = True
                            break

        # Save and compute – also build detailed file list for response (IO-style)
        results = {}
        any_saved = False
        versions_to_process = [forced_version] if forced_version else ["gated", "ungated"]
        detailed_files = {}  # version -> list of file info

        for ver in versions_to_process:
            if ver not in version_groups:
                continue
            cal_path = version_groups[ver]["calendar"]
            sched_path = version_groups[ver]["schedule"]
            persist_dir = os.path.join(DEFAULT_DATA_DIR, "utilization", ver)
            os.makedirs(persist_dir, exist_ok=True)

            # Build detailed file list for this version – IO-style: include original names, types, sizes
            ver_files = []
            for src in version_groups[ver].get("files", []):
                for info in uploaded_file_info:
                    if info["path"] == src:
                        f_type = "calendar" if src == cal_path else "schedule" if src == sched_path else "unknown"
                        ver_files.append({
                            "original": info["original"],
                            "size_kb": info["size_kb"],
                            "size": info["size"],
                            "type": f_type,
                            "field": info["field"],
                            "detected_as": f"{f_type} ({info['original']})",
                        })
                        break
            detailed_files[ver] = ver_files

            # Deep validate files before copying to avoid truncated file causing EOFError later (user reported EOFError)
            if cal_path:
                ok, msg = _is_valid_xlsx_deep(cal_path)
                if not ok:
                    return _json_error(f"Calendar file {os.path.basename(cal_path)} is corrupted/invalid: {msg}. Please re-upload a valid xlsx. Ensure upload completed (file not truncated) and file is .xlsx not .xls. Details: calendar={cal_path}", 400)
                dest = os.path.join(persist_dir, "工作日历快照.xlsx")
                shutil.copyfile(cal_path, dest)
                any_saved = True
            if sched_path:
                ok, msg = _is_valid_xlsx_deep(sched_path)
                if not ok:
                    return _json_error(f"Schedule file {os.path.basename(sched_path)} is corrupted/invalid: {msg}. Please re-upload a valid xlsx. Ensure upload completed (file not truncated). Details: schedule={sched_path}", 400)
                dest = os.path.join(persist_dir, "排产结果表.xlsx")
                shutil.copyfile(sched_path, dest)
                any_saved = True

            cal_persist = os.path.join(persist_dir, "工作日历快照.xlsx")
            sched_persist = os.path.join(persist_dir, "排产结果表.xlsx")
            if os.path.exists(cal_persist) and os.path.exists(sched_persist):
                try:
                    cache = _compute_records(ver, cal_persist, sched_persist)
                    # IO-style detailed message
                    cal_file = next((f for f in ver_files if f['type']=='calendar'), None)
                    sched_file = next((f for f in ver_files if f['type']=='schedule'), None)
                    if cal_file and sched_file:
                        detected_str = f"✅ {ver.capitalize()} Ready: calendar={cal_file['original']}({cal_file['size_kb']}KB) + schedule={sched_file['original']}({sched_file['size_kb']}KB) — {len(cache.lines)} lines, {len(cache.dates)} dates, will be shown in matrix"
                    elif cal_file:
                        detected_str = f"⚠️ {ver.capitalize()} Partial: calendar={cal_file['original']}({cal_file['size_kb']}KB) uploaded, missing schedule — {len(cache.lines)} lines but need both to display"
                    elif sched_file:
                        detected_str = f"⚠️ {ver.capitalize()} Partial: schedule={sched_file['original']}({sched_file['size_kb']}KB) uploaded, missing calendar"
                    else:
                        # Files exist from previous uploads, use persisted names
                        detected_str = f"✅ {ver.capitalize()} Ready: 工作日历快照.xlsx + 排产结果表.xlsx — {len(cache.lines)} lines, {len(cache.dates)} dates (using persisted files)"
                    results[ver] = {
                        "lines": len(cache.lines),
                        "dates": len(cache.dates),
                        "records_shift": len(cache.records_shift),
                        "ready": True,
                        "files": ver_files,
                        "detected": detected_str,
                        "summary": f"{ver}: " + (f"📅 {cal_file['original']} + 📋 {sched_file['original']}" if cal_file and sched_file else "ready")
                    }
                except Exception as ve:
                    import traceback
                    traceback.print_exc()
                    # Check if error is due to corrupted file (EOFError, BadZipFile)
                    err_str = str(ve).lower()
                    is_corrupt = any(kw in err_str for kw in ["eof", "truncated", "badzip", "not a zip", "corrupted", "invalid"])
                    if is_corrupt:
                        # Try to clean up corrupted persist files to allow re-upload
                        try:
                            # Optionally remove corrupted files? Keep for debugging but suggest clear
                            # For safety, don't auto-delete, just mark as error with suggestion
                            pass
                        except:
                            pass
                        results[ver] = {
                            "ready": False,
                            "error": str(ve),
                            "files": ver_files,
                            "detected": f"❌ {ver} failed due to corrupted file: {ve} — Please clear and re-upload. Files may be truncated (network interrupted) or not valid xlsx. Try: 1) Clear (Both Not Ready) 2) Re-upload via zip 3) Ensure files are .xlsx not .xls and upload completes.",
                            "corrupted": True,
                            "suggestion": "Clear and re-upload"
                        }
                    else:
                        results[ver] = {"ready": False, "error": str(ve), "files": ver_files, "detected": f"❌ {ver} failed: {ve}"}
            else:
                has_cal = os.path.exists(cal_persist)
                has_sched = os.path.exists(sched_persist)
                if has_cal or has_sched:
                    cal_file = next((f for f in ver_files if f['type']=='calendar'), None)
                    sched_file = next((f for f in ver_files if f['type']=='schedule'), None)
                    partial_msg = f"Partial upload for {ver}: "
                    if cal_file:
                        partial_msg += f"calendar={cal_file['original']}({cal_file['size_kb']}KB) present, "
                    else:
                        partial_msg += f"calendar={'present' if has_cal else 'missing'}, "
                    if sched_file:
                        partial_msg += f"schedule={sched_file['original']}({sched_file['size_kb']}KB) {'present' if has_sched else 'missing'}"
                    else:
                        partial_msg += f"schedule={'present' if has_sched else 'missing'}"
                    partial_msg += " — upload missing file to complete (need both calendar + schedule)"
                    results[ver] = {
                        "ready": False,
                        "partial": True,
                        "has_calendar": has_cal,
                        "has_schedule": has_sched,
                        "files": ver_files,
                        "message": partial_msg,
                        "detected": partial_msg,
                    }

        if forced_version:
            other = "ungated" if forced_version == "gated" else "gated"
            other_dir = os.path.join(DEFAULT_DATA_DIR, "utilization", other)
            cal_other = os.path.join(other_dir, "工作日历快照.xlsx")
            sched_other = os.path.join(other_dir, "排产结果表.xlsx")
            if os.path.exists(cal_other) or os.path.exists(sched_other):
                has_cal = os.path.exists(cal_other)
                has_sched = os.path.exists(sched_other)
                if has_cal and has_sched:
                    if other not in results:
                        results[other] = {"ready": True, "existing": True, "has_calendar": True, "has_schedule": True, "files": []}
                else:
                    if other not in results:
                        results[other] = {"ready": False, "partial": True, "has_calendar": has_cal, "has_schedule": has_sched, "existing": True, "files": []}

        if not any_saved and not results:
            return _json_error("No valid calendar/schedule files classified. Please upload 工作日历快照.xlsx and 排产结果表.xlsx, or zip containing them.", 400)

        reload_all(DEFAULT_DATA_DIR)

        # Build overall detected message like IO: "✅ Snapshot mode: detected N files — ..." – now with per-version breakdown
        overall_files = []
        ver_summaries = []
        for ver, flist in detailed_files.items():
            if not flist:
                continue
            parts = []
            for f in flist:
                parts.append(f"{f['type']}={f['original']}({f['size_kb']}KB)")
            ver_summaries.append(f"{ver}:[{' + '.join(parts)}]")
            for f in flist:
                overall_files.append(f"{ver}:{f['original']}")

        # IO-style: show specific uploaded files with sizes and types, like IO's "Detected N files — ..."
        if uploaded_file_info:
            file_list_str = ", ".join([f"{info['original']}({info['size_kb']}KB)" for info in uploaded_file_info])
            ver_detail_str = "; ".join(ver_summaries) if ver_summaries else file_list_str
            detected_msg = f"✅ Utilization upload: detected {len(uploaded_file_info)} files — {file_list_str} | breakdown: {ver_detail_str}"
        else:
            detected_msg = "✅ Upload complete — using existing persisted files"

        # Also build per-version file list for frontend persistent badge (IO-like)
        per_version_summary = {}
        for ver in detailed_files:
            flist = detailed_files[ver]
            if flist:
                per_version_summary[ver] = {
                    "count": len(flist),
                    "files": flist,
                    "text": " + ".join([f"{f['type']}:{f['original']}({f['size_kb']}KB)" for f in flist])
                }

        return jsonify({
            "ok": True,
            "results": results,
            "saved_versions": list(results.keys()),
            "files": uploaded_file_info,
            "detailed_files": detailed_files,
            "per_version_summary": per_version_summary,
            "ver_summaries": ver_summaries,
            "detected_message": detected_msg,
            "file_list_detail": file_list_str if uploaded_file_info else "",
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(f"Upload failed: {e}")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        shutil.rmtree(tmp_extract_dir, ignore_errors=True)

@util_bp.route("/api/utilization/demo/load", methods=["POST"])
def api_demo_load():
    """
    Load demo from data/IVY20260721Gated folder as gated version
    """
    try:
        base = pathlib.Path(DEFAULT_DATA_DIR) / "IVY20260721Gated"
        cal = base / "工作日历快照.xlsx"
        sched = base / "排产结果表.xlsx"
        if not cal.exists() or not sched.exists():
            return _json_error(f"Demo files not found in {base}", 404)
        persist_dir = pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / "gated"
        persist_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(str(cal), str(persist_dir / "工作日历快照.xlsx"))
        shutil.copyfile(str(sched), str(persist_dir / "排产结果表.xlsx"))
        reload_all(DEFAULT_DATA_DIR)
        caches = get_all_caches(DEFAULT_DATA_DIR)
        cache = caches.get('gated')
        if not cache:
            return _json_error("Failed to load after copy", 500)
        return jsonify({
            "ok": True,
            "version": "gated",
            "lines": len(cache.lines),
            "dates": len(cache.dates),
            "records_shift": len(cache.records_shift),
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(str(e))


@util_bp.route("/api/utilization/clear", methods=["POST", "DELETE"])
def api_clear():
    """
    Clear cached data for a specific version or all.
    If user doesn't want to see a previously loaded module, clear it so it becomes Not Ready.
    Query param: version=gated|ungated|all (default all)
    """
    try:
        version = (request.args.get('version') or request.form.get('version') or (request.json.get('version') if request.is_json else '') or '').strip().lower()
        if not version:
            version = request.args.get('ver') or ''
            version = version.strip().lower() if version else 'all'

        valid = {'gated', 'ungated', 'all'}
        if version not in valid:
            # If invalid, treat as all for safety, but return error if clearly wrong
            if version == '':
                version = 'all'
            else:
                return _json_error(f"Invalid version param: {version}. Use gated|ungated|all", 400)

        from app.modules.utilization_report.engine import _CACHES

        cleared = []
        if version == 'all':
            to_clear = ['gated', 'ungated']
        else:
            to_clear = [version]

        for ver in to_clear:
            persist_dir = pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / ver
            # Remove cache pickle files and xlsx files
            try:
                if persist_dir.exists():
                    # Remove all files inside, but keep dir for future uploads (or remove dir entirely)
                    for f in persist_dir.glob("*"):
                        try:
                            if f.is_file():
                                f.unlink()
                            elif f.is_dir():
                                shutil.rmtree(str(f))
                        except Exception as e:
                            print(f"[util] clear file failed {f}: {e}")
                    # Optionally remove the directory itself to indicate Not Ready (empty)
                    # Keep empty dir to avoid confusion, but it's okay to leave empty
                    # If you want to remove dir entirely, uncomment:
                    # shutil.rmtree(str(persist_dir), ignore_errors=True)
            except Exception as e:
                print(f"[util] clear dir failed {ver}: {e}")

            # Clear from in-memory cache
            # _CACHES keys are like "base_dir::version"
            keys_to_remove = [k for k in list(_CACHES.keys()) if k.endswith(f"::{ver}")]
            for k in keys_to_remove:
                _CACHES.pop(k, None)
            cleared.append(ver)

        # After clearing, reload remaining caches (if any)
        try:
            reload_all(DEFAULT_DATA_DIR)
        except Exception:
            pass

        # Get current status after clear
        caches = get_all_caches(DEFAULT_DATA_DIR)
        remaining = list(caches.keys())

        return jsonify({
            "ok": True,
            "cleared": cleared,
            "remaining_versions": remaining,
            "message": f"Cleared {', '.join(cleared)}. Remaining: {', '.join(remaining) if remaining else 'none — all Not Ready'}"
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(f"Clear failed: {e}")

# Templates - detailed schema and downloadable files
@util_bp.route("/api/utilization/templates/schema", methods=["GET"])
def api_schema():
    return jsonify({
        "calendar": {
            "file": "工作日历快照.xlsx",
            "required": True,
            "description": "Working calendar snapshot — defines capacity per line per shift",
            "fields": [
                ["LINE_CODE", "String", "Line code", "AL1-PKG", "Required"],
                ["PLAN_DATE", "Date", "Date (YYYY/MM/DD or YYYY-MM-DD)", "2026/6/2", "Required"],
                ["SHIFT_NAME", "String", "Shift name, e.g. 白班 / 夜班 / Day / Night", "白班", "Required"],
                ["PLAN_TYPE", "String", "Type: UPH / 工时 (working hours) / 效率 (efficiency) / 良率", "UPH", "Required"],
                ["PLAN_VALUE", "Number", "Value: UPH=300, 工时=10, 效率=0.9", "320", "Required"],
                ["PLAN_ITEM", "String", "Fixed INPUT — only INPUT rows are used for capacity", "INPUT", "Required"],
            ],
            "formula": "Capacity per shift = UPH × 效率 × 工时 where PLAN_ITEM=INPUT",
            "example": "For line AL1-PKG on 2026-06-02 白班: UPH=300, 效率=0.9, 工时=10 => Capacity=2700",
            "note": "Must have 3 rows per line/date/shift (one per PLAN_TYPE). Missing UPH/效率/工时 defaults to 0 => capacity 0.",
        },
        "schedule": {
            "file": "排产结果表.xlsx",
            "required": True,
            "description": "Schedule result — defines load per line per shift",
            "fields": [
                ["LINE_CODE", "String", "Line", "AL1-PKG", "Required"],
                ["SHIFT_NAME", "String", "Shift", "白班", "Required"],
                ["PLAN_ITEM", "String", "Must be INPUT — only INPUT is summed as load", "INPUT", "Required"],
                ["PLAN_DATE", "Date", "Date", "2026/6/2", "Required"],
                ["PLAN_VALUE", "Number", "Qty", "1000", "Required"],
            ],
            "formula": "Load per shift = sum PLAN_VALUE where PLAN_ITEM=INPUT",
            "example": "If schedule has 1000 INPUT for AL1-PKG 2026-06-02 白班, Load=1000",
            "note": "Utilization = Load / Capacity, capped at 100% for display, overload flagged if >100%",
        },
        "upload_requirements": {
            "gated": "Independent module — can upload only Gated (calendar + schedule) and show Gated only, Ungated can be empty",
            "ungated": "Independent module — can upload only Ungated and show Ungated only",
            "files_per_module": "2 files per module: 1 calendar (工作日历快照.xlsx) + 1 schedule (排产结果表.xlsx), or 1 zip containing both",
            "both_modules": "If both Gated and Ungated uploaded, matrix shows both with version_type column (Gated yellow, Ungated green)",
        }
    })


def _util_xlsx_buf(header, rows=None):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows or []:
        ws.append(r)
    # Add some styling for header? keep simple
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _util_zip_buf(empty=True, use_existing_demo=False):
    """
    Create zip with calendar and schedule
    empty=True -> only headers
    empty=False -> sample data or existing demo files if available
    """
    import io
    import zipfile
    import pathlib

    cal_header = ["LINE_CODE", "PLAN_DATE", "SHIFT_NAME", "PLAN_TYPE", "PLAN_VALUE", "PLAN_ITEM"]
    sched_header = ["LINE_CODE", "SHIFT_NAME", "PLAN_ITEM", "PLAN_DATE", "PLAN_VALUE"]

    if empty:
        cal_rows = []
        sched_rows = []
    else:
        # Sample rows for demo
        cal_rows = [
            ["AL1-PKG", "2026/6/2", "白班", "UPH", 300, "INPUT"],
            ["AL1-PKG", "2026/6/2", "白班", "工时", 10, "INPUT"],
            ["AL1-PKG", "2026/6/2", "白班", "效率", 0.9, "INPUT"],
            ["AL1-PKG", "2026/6/2", "夜班", "UPH", 280, "INPUT"],
            ["AL1-PKG", "2026/6/2", "夜班", "工时", 10, "INPUT"],
            ["AL1-PKG", "2026/6/2", "夜班", "效率", 0.85, "INPUT"],
            ["AL1-FAT", "2026/6/2", "白班", "UPH", 320, "INPUT"],
            ["AL1-FAT", "2026/6/2", "白班", "工时", 10, "INPUT"],
            ["AL1-FAT", "2026/6/2", "白班", "效率", 0.92, "INPUT"],
        ]
        sched_rows = [
            ["AL1-PKG", "白班", "INPUT", "2026/6/2", 1500],
            ["AL1-PKG", "夜班", "INPUT", "2026/6/2", 1200],
            ["AL1-FAT", "白班", "INPUT", "2026/6/2", 2000],
            ["AL1-FAT", "白班", "INPUT", "2026/6/3", 1800],
        ]

    # Try to use existing demo files if use_existing_demo and files exist
    if use_existing_demo and not empty:
        # Look for IVY demo or utilization/gated demo
        candidates = [
            pathlib.Path(DEFAULT_DATA_DIR) / "IVY20260721Gated" / "工作日历快照.xlsx",
            pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / "gated" / "工作日历快照.xlsx",
            pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / "gated" / "工作日历快照.xlsx",
        ]
        # Actually for zip demo we might want to include real files if they exist
        # For simplicity, if real files exist, we will zip those files directly (preserve content)
        real_cal = None
        real_sched = None
        for p in [pathlib.Path(DEFAULT_DATA_DIR) / "IVY20260721Gated" / "工作日历快照.xlsx",
                  pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / "gated" / "工作日历快照.xlsx"]:
            if p.exists():
                real_cal = p
                break
        for p in [pathlib.Path(DEFAULT_DATA_DIR) / "IVY20260721Gated" / "排产结果表.xlsx",
                  pathlib.Path(DEFAULT_DATA_DIR) / "utilization" / "gated" / "排产结果表.xlsx"]:
            if p.exists():
                real_sched = p
                break
        if real_cal and real_sched:
            # Return zip with real files
            zb = io.BytesIO()
            with zipfile.ZipFile(zb, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.write(str(real_cal), arcname="工作日历快照.xlsx")
                zf.write(str(real_sched), arcname="排产结果表.xlsx")
            zb.seek(0)
            return zb

    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("工作日历快照.xlsx", _util_xlsx_buf(cal_header, cal_rows).getvalue())
        zf.writestr("排产结果表.xlsx", _util_xlsx_buf(sched_header, sched_rows).getvalue())
    zb.seek(0)
    return zb


@util_bp.route("/api/utilization/templates/template", methods=["GET"])
def api_template():
    """Download empty templates zip (calendar + schedule with headers only)"""
    from flask import send_file
    return send_file(_util_zip_buf(empty=True), mimetype="application/zip", as_attachment=True, download_name="utilization_templates.zip")


@util_bp.route("/api/utilization/templates/calendar", methods=["GET"])
def api_template_calendar():
    """Download single empty calendar template"""
    from flask import send_file
    header = ["LINE_CODE", "PLAN_DATE", "SHIFT_NAME", "PLAN_TYPE", "PLAN_VALUE", "PLAN_ITEM"]
    return send_file(_util_xlsx_buf(header, []), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="工作日历快照_template.xlsx")


@util_bp.route("/api/utilization/templates/schedule", methods=["GET"])
def api_template_schedule():
    """Download single empty schedule template"""
    from flask import send_file
    header = ["LINE_CODE", "SHIFT_NAME", "PLAN_ITEM", "PLAN_DATE", "PLAN_VALUE"]
    return send_file(_util_xlsx_buf(header, []), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="排产结果表_template.xlsx")


@util_bp.route("/api/utilization/templates/demo", methods=["GET"])
def api_demo_download():
    """Download demo data zip - like template download, for Gated module (independent)"""
    from flask import send_file
    return send_file(_util_zip_buf(empty=False, use_existing_demo=True), mimetype="application/zip", as_attachment=True, download_name="utilization_demo_gated.zip")


@util_bp.route("/api/utilization/templates/zip", methods=["GET"])
@util_bp.route("/api/utilization/templates/input_template.xlsx", methods=["GET"])
def api_combined_template():
    """Combined template workbook with 2 sheets: calendar and schedule"""
    import io
    import openpyxl
    from flask import send_file
    wb = openpyxl.Workbook()
    ws_cal = wb.active
    ws_cal.title = "Calendar (工作日历快照)"
    ws_cal.append(["LINE_CODE", "PLAN_DATE", "SHIFT_NAME", "PLAN_TYPE", "PLAN_VALUE", "PLAN_ITEM"])
    ws_cal.append(["AL1-PKG", "2026/6/2", "白班", "UPH", 300, "INPUT"])
    ws_cal.append(["AL1-PKG", "2026/6/2", "白班", "工时", 10, "INPUT"])
    ws_cal.append(["AL1-PKG", "2026/6/2", "白班", "效率", 0.9, "INPUT"])
    ws_sched = wb.create_sheet("Schedule (排产结果表)")
    ws_sched.append(["LINE_CODE", "SHIFT_NAME", "PLAN_ITEM", "PLAN_DATE", "PLAN_VALUE"])
    ws_sched.append(["AL1-PKG", "白班", "INPUT", "2026/6/2", 1500])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="utilization_template.xlsx")


def _build_util_pivot_for_export(caches, mode):
    """Build pivot for export static with flexible filtering preserved"""
    try:
        from app.modules.utilization_report.engine import build_report
        from collections import defaultdict
        all_recs = []
        for ver in caches.keys():
            cache = caches.get(ver)
            if not cache:
                continue
            recs = build_report(cache, mode=mode, line_filter='', date_from='', date_to='')
            for r in recs:
                nr = dict(r)
                nr['version_type'] = ver.capitalize()
                nr['version_key'] = ver
                all_recs.append(nr)

        if not all_recs:
            return {"mode": mode, "columns": [], "rows": [], "detail": {}, "lines": [], "total_lines": 0, "total_cols": 0}

        if mode == "shift":
            cols_set = set()
            for r in all_recs:
                col = f"{r['plan_date']}|{r['shift_name']}"
                cols_set.add(col)
            cols = sorted(list(cols_set))
        else:
            cols = sorted(list(set(r['plan_date'] for r in all_recs)))

        max_cols = 1000
        if len(cols) > max_cols:
            cols = cols[:max_cols]

        matrix = {}
        detail = {}
        lines_set = set()
        for r in all_recs:
            key = (r['line_code'], r['version_type'])
            lines_set.add(r['line_code'])
            if key not in matrix:
                matrix[key] = {}
                detail[key] = {}
            col = f"{r['plan_date']}|{r['shift_name']}" if mode == "shift" else r['plan_date']
            if col not in cols:
                continue
            util_pct = r.get('utilization_pct', 0)
            capped = min(100.0, util_pct) if util_pct is not None else 0
            matrix[key][col] = capped
            detail[key][col] = {
                'util_raw': util_pct,
                'util_capped': capped,
                'load': r['load'],
                'capacity': r['capacity'],
                'uph': r['uph'],
                'efficiency': r['efficiency'],
                'working_hours': r['working_hours'],
                'is_overload': r.get('is_overload', False),
            }

        def sort_key(k):
            line, vtype = k
            order = 0 if vtype.lower() == 'gated' else 1
            return (line, order, vtype)
        sorted_keys = sorted(matrix.keys(), key=sort_key)

        rows = []
        last_line = None
        for (line, vtype) in sorted_keys:
            row = {
                'line_code': line,
                'version_type': vtype,
                'is_new_line': line != last_line,
            }
            last_line = line
            for col in cols:
                row[col] = matrix[(line, vtype)].get(col, None)
            rows.append(row)

        detail_json = {}
        for (line, vtype), col_map in detail.items():
            key_str = f"{line}||{vtype}"
            detail_json[key_str] = col_map

        return {
            "mode": mode,
            "versions": list(caches.keys()),
            "columns": cols,
            "rows": rows,
            "detail": detail_json,
            "lines": sorted(list(lines_set)),
            "total_lines": len(sorted_keys),
            "total_cols": len(cols),
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"mode": mode, "columns": [], "rows": [], "detail": {}, "lines": [], "total_lines": 0, "total_cols": 0, "error": str(e)}


@util_bp.route("/api/utilization/export/static", methods=["GET"])
def api_export_static():
    """Export current utilization dashboard as static HTML with embedded data and flexible filtering (like campus-planning-system/frontend/dist)
    Called when user clicks Download Static HTML button — triggers export_static logic for current data
    """
    try:
        import json
        import datetime
        caches = get_all_caches(DEFAULT_DATA_DIR)
        if not caches:
            return _json_error("No data to export. Upload Gated or Ungated first.", 404)

        # Build status
        from app.modules.utilization_report.engine import _CACHES
        cached_versions = {}
        for k in _CACHES.keys():
            if k.startswith(DEFAULT_DATA_DIR + "::"):
                ver = k.split("::")[-1]
                c = _CACHES[k]
                cached_versions[ver] = {
                    'lines': len(c.lines),
                    'dates': len(c.dates),
                    'shifts': len(c.shifts),
                    'records_shift': len(c.records_shift),
                    'records_day': len(c.records_day),
                }

        status = {
            "loaded": True,
            "versions": list(caches.keys()),
            "details": cached_versions,
            "files_found": list(caches.keys()),
        }
        # Build meta
        primary = caches.get('gated') or list(caches.values())[0]
        meta = {
            "lines": primary.lines,
            "dates": primary.dates,
            "shifts": primary.shifts,
            "versions": list(caches.keys()),
            "date_min": min(primary.dates) if primary.dates else "",
            "date_max": max(primary.dates) if primary.dates else "",
        }

        # Build pivots for day and shift
        pivot_day = _build_util_pivot_for_export(caches, mode='day')
        pivot_shift = _build_util_pivot_for_export(caches, mode='shift')

        # Build static HTML with embedded data and flexible filtering (full format preserved)
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        static_data = {
            "status": status,
            "meta": meta,
            "pivotDay": pivot_day,
            "pivotShift": pivot_shift,
            "currentMode": "day"
        }
        # Use json dumps with ensure_ascii=False, escape < for safety
        import html as html_lib
        static_json = json.dumps(static_data, ensure_ascii=False, default=str).replace("<", "\\u003c")

        html_content = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Line Utilization - Static BI - {now_str}</title>
<style>
body{{font-family:Arial,sans-serif;margin:16px;background:#f8fafc;color:#1e293b}}
h1{{font-size:18px;margin-bottom:4px}}
.sub{{font-size:11px;color:#64748b;margin-bottom:10px}}
.status{{margin:10px 0;padding:10px;background:#fff;border:1px solid #e2e8f0;border-radius:6px;font-size:12px}}
.filters{{background:#fff;border:1px solid #e2e8f0;border-radius:6px;padding:10px;margin:10px 0;display:flex;flex-wrap:wrap;gap:12px;align-items:flex-end}}
.filter-group{{display:flex;flex-direction:column;gap:4px;min-width:160px}}
.filter-group label{{font-size:10px;font-weight:600;color:#64748b;text-transform:uppercase}}
.filter-group input, .filter-group select{{padding:5px 8px;border:1px solid #cbd5e1;border-radius:5px;font-size:12px}}
.btn{{padding:5px 12px;border:1px solid #3b82f6;background:#3b82f6;color:#fff;border-radius:5px;font-size:12px;cursor:pointer}}
.btn-outline{{background:#fff;color:#64748b;border-color:#cbd5e1}}
.table-wrapper{{overflow:auto;max-height:80vh;border:1px solid #e2e8f0;border-radius:6px;background:#fff;margin-top:10px}}
table{{border-collapse:collapse;font-size:12px;white-space:nowrap;width:max-content;min-width:100%}}
th{{background:#1e293b;color:#fff;padding:6px 8px;position:sticky;top:0;z-index:2;border-right:1px solid #334155}}
td{{padding:4px 6px;border-bottom:1px solid #e2e8f0;border-right:1px solid #f1f5f9;text-align:center;min-width:68px}}
td.frozen{{position:sticky;left:0;background:#fff;z-index:1;min-width:68px;text-align:left;font-weight:500}}
td.frozen.divider-col{{background:#475569 !important;width:5px;min-width:5px;max-width:5px;padding:0 !important}}
th.frozen{{left:0;z-index:3;background:#1e293b}}
th.divider-col{{background:#475569 !important;width:5px;min-width:5px;max-width:5px}}
.util-cell-zero{{color:#cbd5e1}}
.util-cell-red{{background:#fef2f2;color:#991b1b}}
.util-cell-yellow{{background:#fffbeb;color:#92400e}}
.util-cell-green{{background:#ecfdf5;color:#065f46;font-weight:600}}
.type-Gated{{background:#fef3c7;color:#92400e;padding:1px 6px;border-radius:4px;font-size:11px}}
.type-Ungated{{background:#d1fae5;color:#065f46;padding:1px 6px;border-radius:4px;font-size:11px}}
.badge{{display:inline-block;padding:2px 8px;border-radius:10px;font-size:11px;margin-right:4px}}
.badge-ready{{background:#dcfce7;color:#065f46;border:1px solid #86efac}}
.badge-notready{{background:#fef2f2;color:#991b1b;border:1px solid #fecaca}}
</style></head><body>
<h1>⚙️ Line Utilization — Static BI Report (Full Format & Filtering Preserved)</h1>
<div class="sub">Generated: {now_str} | Export via /api/utilization/export/static (calls export_static logic) | Day cols: {len(pivot_day.get('columns',[]))}, Shift cols: {len(pivot_shift.get('columns',[]))}, Lines: {len(meta.get('lines',[]))}</div>
<div class="status" id="static-status"></div>
<div class="filters">
  <div class="filter-group"><label>Version Type</label>
    <div><label><input type="checkbox" id="f-gated" checked> Gated</label> <label><input type="checkbox" id="f-ungated" checked> Ungated</label></div>
  </div>
  <div class="filter-group"><label>Line (comma separated)</label><input type="text" id="f-line" placeholder="e.g. AL1-PKG,AL1-FAT"></div>
  <div class="filter-group"><label>Date From</label><input type="date" id="f-from"></div>
  <div class="filter-group"><label>Date To</label><input type="date" id="f-to"></div>
  <div class="filter-group"><label>Mode</label><div><button class="btn" id="f-mode-day">Day</button> <button class="btn btn-outline" id="f-mode-shift">Shift</button></div></div>
  <div class="filter-group"><label>&nbsp;</label><div><button class="btn" id="f-apply">Apply Filters</button> <button class="btn btn-outline" id="f-clear">Clear</button></div></div>
</div>
<div style="font-size:11px;color:#64748b">Formula: Capacity=UPH×Eff×WH | Load=Σ INPUT | Util%=Load/Capacity capped at 100% | Thick border per Line | Gated yellow, Ungated green</div>
<div id="static-badge" style="margin:8px 0;font-size:11px"></div>
<div class="table-wrapper" id="static-wrapper"><div style="text-align:center;padding:30px;color:#94a3b8">Loading...</div></div>
<div style="margin-top:12px;font-size:10px;color:#94a3b8">Static BI report from Line Utilization dashboard via export_static. All format and filtering preserved like campus-planning-system/frontend/dist. Data embedded at export time.</div>
<script>
const STATIC_DATA = {static_json};

function esc(s){{ return s ? String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;') : ''; }}

let currentMode = STATIC_DATA.currentMode || 'day';
let selectedVersions = new Set(['Gated','Ungated']);
let filterLine = '';
let dateFrom = '';
let dateTo = '';

function getPivot(){{ return currentMode==='day' ? STATIC_DATA.pivotDay : STATIC_DATA.pivotShift; }}

function renderStatus(){{
  const st = STATIC_DATA.status;
  const versions = (st && st.versions) ? st.versions : [];
  const gatedReady = versions.includes('gated');
  const ungatedReady = versions.includes('ungated');
  let html = '';
  if(gatedReady && ungatedReady) html = '<span class="badge badge-ready">✅ Gated: Ready</span><span class="badge badge-ready">✅ Ungated: Ready</span> — Showing both';
  else if(gatedReady) html = '<span class="badge badge-ready">✅ Gated: Ready</span><span class="badge badge-notready">❌ Ungated: Not Ready (empty)</span> — Showing Gated only';
  else if(ungatedReady) html = '<span class="badge badge-notready">❌ Gated: Not Ready</span><span class="badge badge-ready">✅ Ungated: Ready</span> — Showing Ungated only';
  else html = '<span class="badge badge-notready">❌ Gated: Not Ready</span><span class="badge badge-notready">❌ Ungated: Not Ready</span> — No data';
  document.getElementById('static-status').innerHTML = html;
}}

function renderMatrix(){{
  const pivot = getPivot();
  const cols = pivot.columns || [];
  const rows = pivot.rows || [];
  const detail = pivot.detail || {{}};
  const wrapper = document.getElementById('static-wrapper');
  if(!rows || rows.length===0){{
    wrapper.innerHTML = '<div style="text-align:center;padding:30px;color:#991b1b;background:#fef2f2;border:1px solid #fecaca;border-radius:6px">No data — both modules Not Ready or filtered out</div>';
    document.getElementById('static-badge').textContent = '0 rows';
    return;
  }}
  let filteredRows = rows.filter(r=>{{
    const v = (r.version_type||'').toLowerCase();
    if(v==='gated' && !selectedVersions.has('Gated')) return false;
    if(v==='ungated' && !selectedVersions.has('Ungated')) return false;
    if(filterLine){{
      const lineFilters = filterLine.toLowerCase().split(',').map(s=>s.trim()).filter(Boolean);
      if(lineFilters.length>0){{
        const lc = (r.line_code||'').toLowerCase();
        if(!lineFilters.some(f=> lc.includes(f))) return false;
      }}
    }}
    return true;
  }});
  let filteredCols = cols;
  if(dateFrom || dateTo){{
    filteredCols = cols.filter(c=>{{
      let d = c;
      if(c.includes('|')) d = c.split('|')[0];
      if(dateFrom && d < dateFrom) return false;
      if(dateTo && d > dateTo) return false;
      return true;
    }});
  }}

  const frozenCols = [{{key:'line_code',label:'Line',width:130}},{{key:'version_type',label:'Version Type',width:110}}];
  let left=0; frozenCols.forEach(c=>{{ c._left=left; left+=c.width; }});
  const dividerLeft = left;
  let thead = '<tr>';
  frozenCols.forEach(c=>{{ thead += '<th class="frozen" style="left:'+c._left+'px;min-width:'+c.width+'px">'+esc(c.label)+'</th>'; }});
  thead += '<th class="frozen divider-col" style="left:'+dividerLeft+'px;min-width:5px"></th>';
  filteredCols.forEach(col=>{{
    let label = col;
    let sub='';
    if(col.includes('|')){{ const parts=col.split('|'); label=parts[0]; sub=parts[1]; try{{ const d=new Date(label); if(!isNaN(d)) label=(d.getMonth()+1)+'/'+d.getDate(); }}catch(e){{}} }}else{{ try{{ const d=new Date(col); if(!isNaN(d)) label=(d.getMonth()+1)+'/'+d.getDate(); }}catch(e){{}} }}
    thead += '<th style="min-width:68px" title="'+esc(col)+'">'+esc(label)+(sub?'<br><span style="font-size:9px;color:#cbd5e1">'+esc(sub)+'</span>':'')+'</th>';
  }});
  thead += '</tr>';

  let tbody='';
  let lastLine=null;
  filteredRows.forEach(r=>{{
    const isNewLine = r.line_code !== lastLine;
    lastLine = r.line_code;
    const vType = r.version_type||'';
    tbody += '<tr class="'+(isNewLine?'row-new-line':'')+'">';
    frozenCols.forEach(c=>{{
      const isLast = c===frozenCols[frozenCols.length-1];
      const extra = isLast ? ' frozen-last' : '';
      let val='';
      if(c.key==='line_code') val=esc(r.line_code);
      else if(c.key==='version_type') val='<span class="type-'+esc(vType)+'">'+esc(vType)+'</span>';
      tbody += '<td class="frozen data-cell'+extra+'" style="left:'+c._left+'px;min-width:'+c.width+'px">'+val+'</td>';
    }});
    tbody += '<td class="divider-col frozen" style="left:'+dividerLeft+'px"></td>';
    filteredCols.forEach(col=>{{
      const keyStr = r.line_code+'||'+vType;
      const cellDetail = detail[keyStr] && detail[keyStr][col];
      const cellVal = r[col];
      if(cellVal==null){{ tbody += '<td class="data-cell" style="background:#f8fafc"></td>'; }}
      else{{
        let cls='';
        if(cellVal===0) cls='util-cell-zero';
        else if(cellVal<60) cls='util-cell-red';
        else if(cellVal<80) cls='util-cell-yellow';
        else cls='util-cell-green';
        tbody += '<td class="data-cell '+cls+'">'+Math.round(cellVal)+'%</td>';
      }}
    }});
    tbody += '</tr>';
  }});

  wrapper.innerHTML = '<table><thead>'+thead+'</thead><tbody>'+tbody+'</tbody></table>';
  document.getElementById('static-badge').textContent = filteredRows.length+' rows × '+filteredCols.length+' cols (filtered from '+rows.length+' rows × '+cols.length+' cols) | Thick border per Line';
}}

document.getElementById('f-gated').addEventListener('change', (e)=>{{ if(e.target.checked) selectedVersions.add('Gated'); else selectedVersions.delete('Gated'); }});
document.getElementById('f-ungated').addEventListener('change', (e)=>{{ if(e.target.checked) selectedVersions.add('Ungated'); else selectedVersions.delete('Ungated'); }});
document.getElementById('f-line').addEventListener('input', (e)=>{{ filterLine = e.target.value; }});
document.getElementById('f-from').addEventListener('change', (e)=>{{ dateFrom = e.target.value; }});
document.getElementById('f-to').addEventListener('change', (e)=>{{ dateTo = e.target.value; }});
document.getElementById('f-mode-day').addEventListener('click', ()=>{{ currentMode='day'; document.getElementById('f-mode-day').className='btn'; document.getElementById('f-mode-shift').className='btn btn-outline'; renderMatrix(); }});
document.getElementById('f-mode-shift').addEventListener('click', ()=>{{ currentMode='shift'; document.getElementById('f-mode-shift').className='btn'; document.getElementById('f-mode-day').className='btn btn-outline'; renderMatrix(); }});
document.getElementById('f-apply').addEventListener('click', renderMatrix);
document.getElementById('f-clear').addEventListener('click', ()=>{{ selectedVersions=new Set(['Gated','Ungated']); filterLine=''; dateFrom=''; dateTo=''; document.getElementById('f-gated').checked=true; document.getElementById('f-ungated').checked=true; document.getElementById('f-line').value=''; document.getElementById('f-from').value=''; document.getElementById('f-to').value=''; currentMode='day'; renderMatrix(); }});

renderStatus();
renderMatrix();
</script>
</body></html>"""
        from flask import Response
        return Response(html_content, mimetype='text/html', headers={
            "Content-Disposition": f"attachment; filename=utilization_static_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return _json_error(f"Export static failed: {e}")
