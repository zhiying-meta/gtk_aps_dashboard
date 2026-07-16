"""
Generate single xlsx template with 6 sheets + schema JSON for frontend
"""
import openpyxl, os, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")

HDR_Fill = PatternFill("solid", fgColor="1E293B")
HDR_Font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
C_Font = Font(name="微软雅黑", size=10)
THIN = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))

SCHEMA = {}

def add_sheet(wb, sheet_name, headers, schema_fields, schema_note):
    ws = wb.create_sheet(title=sheet_name)
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = HDR_Font; c.fill = HDR_Fill; c.border = THIN
        c.alignment = Alignment(horizontal='center', vertical='center')
    SCHEMA[sheet_name] = {"fields": schema_fields, "note": schema_note}
    print(f"  Sheet: {sheet_name} ({len(headers)} cols)")

# Create template workbook
wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

add_sheet(wb, "sku_master",
    headers=["SKU","Usage","Style","Color","GB_PN","FR_PN","LT_PN","RT_PN","Pallet_Qty"],
    schema_fields=[
        ("SKU","String","Part Number","SK-1001879-01"),
        ("Usage","String","Usage: MP/Dummy/Demo","MP"),
        ("Style","String","Frame Style","Rectangle M"),
        ("Color","String","Color (UPPERCASE)","BLACK"),
        ("GB_PN","String","GB-level Part Number","GB-Rec M-BLACK"),
        ("FR_PN","String","Frame Part Number","FR-Rec M-BLACK"),
        ("LT_PN","String","Left Temple Part Number","LT-Rec M-BLACK"),
        ("RT_PN","String","Right Temple Part Number","RT-Rec M-BLACK"),
        ("Pallet_Qty","Integer","Pallet Quantity (default: 864)","864"),
    ],
    schema_note="SKU master table (required). One row per SKU. Default Pallet_Qty: 864.")

add_sheet(wb, "plan_output_gated",
    headers=["PN","2026-04-16","2026-04-17","2026-04-18","..."],
    schema_fields=[
        ("PN","String","Part Number (SKU or GB/FR/LT/RT PN)","SK-1001879-01"),
        ("2026-04-16","Integer","Daily OUTPUT quantity","0"),
        ("...","...","Subsequent columns: one per day, format yyyy-MM-dd","..."),
    ],
    schema_note="Gated daily production output. First column: PN, subsequent columns: dates (yyyy-MM-dd), values: daily output. Unlimited date columns.")

add_sheet(wb, "plan_output_ungated",
    headers=["PN","2026-04-16","2026-04-17","2026-04-18","..."],
    schema_fields=[
        ("PN","String","Part Number (SKU or GB/FR/LT/RT PN)","SK-1001879-01"),
        ("2026-04-16","Integer","Daily OUTPUT quantity","0"),
        ("...","...","Subsequent columns: one per day","..."),
    ],
    schema_note="Ungated daily production output. Same format as plan_output_gated.")

add_sheet(wb, "forecast",
    headers=["SKU","2026-04-04","2026-04-11","2026-04-18","..."],
    schema_fields=[
        ("SKU","String","Part Number (SKU level only)","SK-1001879-01"),
        ("2026-04-04","Integer","Weekly Forecast (week ending Saturday)","0"),
        ("...","...","Subsequent columns: Saturday dates","..."),
    ],
    schema_note="Weekly Forecast data. Column headers: Saturday dates, values: weekly forecast quantity.")

add_sheet(wb, "ctb_sku_cum",
    headers=["SKU","2026-02-04","2026-02-05","2026-02-06","..."],
    schema_fields=[
        ("SKU","String","Part Number (SKU level)","SK-1001879-01"),
        ("2026-02-04","Integer","Daily CTB cumulative value","0"),
        ("...","...","Subsequent columns: one per day","..."),
    ],
    schema_note="SKU-level CTB cumulative. Daily data, values are cumulative CTB up to that date (not daily increment).")

add_sheet(wb, "ctb_gb_cum",
    headers=["PN","2026-02-04","2026-02-05","2026-02-06","..."],
    schema_fields=[
        ("PN","String","GB Part Number","GB-Rec M-BLACK"),
        ("2026-02-04","Integer","Daily GB CTB cumulative value","0"),
        ("...","...","Subsequent columns: one per day","..."),
    ],
    schema_note="GB-level CTB cumulative. Optional. Leave empty if no GB-level CTB data.")

template_path = os.path.join(TEMPLATE_DIR, "input_template.xlsx")
wb.save(template_path)
print(f"\nSaved: {template_path}")

# Schema JSON
with open(os.path.join(TEMPLATE_DIR, "schema.json"), "w", encoding="utf-8") as f:
    json.dump(SCHEMA, f, ensure_ascii=False, indent=2)
print("Schema JSON saved")
