"""
Plan merge engine: process uploaded xlsx data → report rows + generate formatted Excel
"""
import io
from datetime import datetime, timedelta
from collections import defaultdict
from functools import lru_cache

from app.modules.plan_merge.config import DEFAULT_PALLET_QTY

_DATE_FORMATS = [
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%Y年%m月%d日",
    "%Y年%m月%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%Y%m%d",
]

@lru_cache(maxsize=2048)
def _to_dt(s):
    """Parse date string with multiple format support (cached)."""
    s = str(s).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"unrecognized date: {s}")


@lru_cache(maxsize=2048)
def _to_saturday_label(ds):
    try: dt = _to_dt(ds)
    except: return ds
    dow = dt.weekday()
    sat = dt + timedelta(days=5 - dow)
    return sat.strftime("%Y-%m-%d")

# cache for week label per (ds, cut_day)
@lru_cache(maxsize=4096)
def _date_to_week_label_cached(ds, cut_day):
    td_map = {"Monday":0,"Tuesday":1,"Wednesday":2,"Thursday":3,"Friday":4,"Saturday":5,"Sunday":6}
    td = td_map.get(cut_day, 5)
    dt = _to_dt(ds)
    cd = dt.weekday()
    diff = (td - cd) % 7
    week_end = dt + timedelta(days=diff)
    return _to_saturday_label(week_end.strftime("%Y-%m-%d"))


def aggregate_cumulative(daily, cut_day):
    """Compute cumulative sum up to each cut_day (cached)."""
    date_list = sorted(daily.keys())
    if not date_list:
        return {}

    def date_to_week_label(ds):
        return _date_to_week_label_cached(ds, cut_day)

    first_wl = date_to_week_label(date_list[0])
    last_wl = date_to_week_label(date_list[-1])

    weeks = defaultdict(list)
    for ds in date_list:
        wl = date_to_week_label(ds)
        weeks[wl].append(ds)

    all_weeks = []
    cur = _to_dt(first_wl)
    end = _to_dt(last_wl)
    while cur <= end:
        all_weeks.append(cur.strftime("%Y-%m-%d"))
        cur += timedelta(days=7)

    result = {}
    running = 0.0
    for wl in all_weeks:
        for ds in weeks.get(wl, []):
            running += daily.get(ds, 0)
        if running > 0:
            result[wl] = round(running, 0)
    return result


def aggregate_etd_from_packout(daily, etd_cut, offset_days):
    """
    ETD cumulative uses Packout cumulative from offset_days before.

    Requirement: if n=2, ETD on date D should use Packout cum on D-n.
    Physically: Packout daily at D becomes ETD at D+n.
    Implementation: shift daily output forward by offset_days, then aggregate.

    e.g. Packout 100 on 2024-01-08 with offset 2 → ETD daily 100 on 2024-01-10.
    Thus ETD_cum[2024-01-10] includes that 100 via P_cum[2024-01-08].
    """
    if not daily:
        return {}
    try:
        off = int(offset_days)
    except Exception:
        off = 0
    if off <= 0:
        return aggregate_cumulative(daily, etd_cut)

    shifted = {}
    for ds, qty in daily.items():
        try:
            dt = _to_dt(ds)
        except Exception:
            continue
        new_dt = dt + timedelta(days=off)
        new_ds = new_dt.strftime("%Y-%m-%d")
        shifted[new_ds] = shifted.get(new_ds, 0.0) + float(qty)
    return aggregate_cumulative(shifted, etd_cut)


def extract_weekly_cum(daily, cut_day):
    """For already-cumulative data (CTB): take the value at each week end."""
    dow_map = {"Monday":0,"Tuesday":1,"Wednesday":2,"Thursday":3,"Friday":4,"Saturday":5,"Sunday":6}
    td = dow_map.get(cut_day, 5)
    date_list = sorted(daily.keys())
    weeks = defaultdict(list)
    for ds in date_list:
        try: dt = _to_dt(ds)
        except: continue
        cd = dt.weekday()
        diff = (td - cd) % 7
        week_end = dt + timedelta(days=diff)
        sat_label = _to_saturday_label(week_end.strftime("%Y-%m-%d"))
        weeks[sat_label].append(ds)
    result = {}
    for wl in sorted(weeks.keys()):
        dl = weeks[wl]
        dl.sort()
        if dl and daily.get(dl[-1]):
            result[wl] = round(daily[dl[-1]], 0)
    return result


def process_uploaded_data(file_map, config):
    from app.modules.plan_merge.utils import read_uploaded_xlsx, read_sku_master_from_ws
    from app.modules.plan_merge.config import DEFAULT_ETD_PACKOUT_OFFSET

    def _parse_offset(v, fallback):
        try:
            if v is None or v == "":
                return fallback
            return max(0, int(v))
        except Exception:
            return fallback

    cfg = {
        "exf_cut": config.get("exf_cut", "Saturday"),
        "etd_cut": config.get("etd_cut", "Saturday"),
        "output_cut": config.get("output_cut", "Wednesday"),
        "gb_cut": config.get("gb_cut", "Tuesday"),
        "etd_packout_offset": _parse_offset(config.get("etd_packout_offset"), DEFAULT_ETD_PACKOUT_OFFSET),
    }

    fp = file_map.get("main") or file_map.get("sku") or next(iter(file_map.values()), None)
    if not fp:
        raise ValueError("No file uploaded")

    sheets, missing_sheets = read_uploaded_xlsx(fp)

    if "sku" not in sheets:
        raise ValueError("Missing required sheet: sku_master")

    sku_ws = sheets["sku"]
    sku_attrs, sku_to_gb, sku_pallet, gb_style_color = read_sku_master_from_ws(sku_ws)
    all_skus = set(sku_attrs.keys())

    plan_gated = sheets.get("gated", {})
    plan_ungated = sheets.get("ungated", {})
    plan_fcst = sheets.get("fcst", {})
    ctb_sku = sheets.get("ctb", {})
    ctb_gb = sheets.get("ctb_gb", {})

    all_pns = set(sku_attrs.keys())
    for d in [plan_gated, plan_ungated, plan_fcst, ctb_sku]:
        all_pns.update(k for k in d if k in sku_attrs)
    all_skus = sorted(s for s in all_pns if s in sku_attrs)

    agg = {}
    # Packout weekly cum (no offset)
    for label, data, cut in [("GATED_PACK", plan_gated, cfg["output_cut"]),
                              ("UNGATED_PACK", plan_ungated, cfg["output_cut"]),
                              ("FCST", plan_fcst, cfg["exf_cut"])]:
        agg[label] = {}
        for sku in all_skus:
            daily = data.get(sku, {})
            agg[label][sku] = aggregate_cumulative(daily, cut)

    # ETD uses Packout cum with offset: ETD(D) = Packout_cum(D - offset)
    # Implemented by shifting daily forward by offset then aggregating to ETD cut
    offset_n = cfg.get("etd_packout_offset", 2)
    for label, data, cut in [("GATED_ETD", plan_gated, cfg["etd_cut"]),
                              ("UNGATED_ETD", plan_ungated, cfg["etd_cut"])]:
        agg[label] = {}
        for sku in all_skus:
            daily = data.get(sku, {})
            agg[label][sku] = aggregate_etd_from_packout(daily, cut, offset_n)

    all_weeks = set()
    for a in agg.values():
        for v in a.values():
            all_weeks.update(v.keys())
    for sku in all_skus:
        if sku in ctb_sku:
            w = extract_weekly_cum(ctb_sku[sku], cfg["etd_cut"])
            all_weeks.update(w.keys())

    # Expand week range to cover earliest raw date across ALL input sheets
    all_raw_dates = set()
    for d in [plan_gated, plan_ungated, plan_fcst, ctb_sku]:
        for v in d.values():
            all_raw_dates.update(v.keys())
    for v in ctb_gb.values():
        all_raw_dates.update(v.keys())
    if all_raw_dates:
        earliest_sat = _to_saturday_label(min(all_raw_dates))
        if not all_weeks or earliest_sat < min(all_weeks):
            end = _to_dt(min(all_weeks)) if all_weeks else _to_dt(earliest_sat)
            cur = _to_dt(earliest_sat)
            extra = []
            while cur < end:
                extra.append(cur.strftime("%Y-%m-%d"))
                cur += timedelta(days=7)
            all_weeks = set(extra) | all_weeks

    all_weeks = sorted(all_weeks)

    def fill(vals): return {w: vals.get(w, None) for w in all_weeks}
    def diff(b, s):
        if not b:
            return {}
        ks = set(list(b.keys()) + list(s.keys()))
        return {k: ((b.get(k) or 0) - (s.get(k) or 0)) for k in ks}

    rows = []
    for sku in all_skus:
        a = sku_attrs.get(sku, {})
        style, color, usage = a.get("Style",""), a.get("Color",""), a.get("Usage","")
        gb = sku_to_gb.get(sku, "")
        pallet = sku_pallet.get(sku, DEFAULT_PALLET_QTY)

        ue = agg["UNGATED_ETD"].get(sku, {})
        up = agg["UNGATED_PACK"].get(sku, {})
        ge = agg["GATED_ETD"].get(sku, {})
        gp = agg["GATED_PACK"].get(sku, {})
        exf = agg["FCST"].get(sku, {})

        ue = {w: (int(v)//pallet)*pallet for w,v in ue.items() if v and v > 0}
        ge = {w: (int(v)//pallet)*pallet for w,v in ge.items() if v and v > 0}

        ctb = {}
        if sku in ctb_sku:
            ctb = extract_weekly_cum(ctb_sku[sku], cfg["etd_cut"])

        base = {"PN": sku, "Usage": usage, "Style": style, "Color": color,
                "GB_PN": gb, "Pallet_Qty": pallet, "_dim": "FG"}
        rows.append({**base, "Version-Type": "ExF", "Version-Detail": "", "Cut Day": cfg["exf_cut"], **fill(exf)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "ETD", "Cut Day": cfg["etd_cut"], **fill(ue)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "ETD vs ExF", "Cut Day": cfg["etd_cut"], **fill(diff(ue, exf))})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "Packout", "Cut Day": cfg["output_cut"], **fill(up)})
        rows.append({**base, "Version-Type": "Ungated", "Version-Detail": "Packout vs ExF", "Cut Day": cfg["output_cut"], **fill(diff(up, exf))})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "ETD", "Cut Day": cfg["etd_cut"], **fill(ge)})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "ETD vs ExF", "Cut Day": cfg["etd_cut"], **fill(diff(ge, exf))})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "Packout", "Cut Day": cfg["output_cut"], **fill(gp)})
        rows.append({**base, "Version-Type": "Gated", "Version-Detail": "Packout vs ExF", "Cut Day": cfg["output_cut"], **fill(diff(gp, exf))})
        rows.append({**base, "Version-Type": "CTB", "Version-Detail": "", "Cut Day": "", **fill(ctb)})

    # ---- Build canonical GB mapping (SKU master is source of truth, case-insensitive) ----
    # Map lower -> canonical SKU master GB
    sku_gb_lower_to_canonical = {}
    for g in sku_to_gb.values():
        if g:
            low = g.lower()
            # Keep first occurrence as canonical (SKU master naming)
            if low not in sku_gb_lower_to_canonical:
                sku_gb_lower_to_canonical[low] = g
    # Also include gb_style_color keys (same set but ensure)
    for g in gb_style_color.keys():
        low = g.lower()
        if low not in sku_gb_lower_to_canonical:
            sku_gb_lower_to_canonical[low] = g

    def to_canonical_gb(gb_pn):
        if not gb_pn:
            return gb_pn
        low = gb_pn.lower()
        if low in sku_gb_lower_to_canonical:
            return sku_gb_lower_to_canonical[low]
        # Alias handling: try to find SKU master GB with same Style+Color (for DEEP BLACK vs Black (low cost))
        # Parse GB-Style-Color
        rest = gb_pn[3:] if gb_pn.upper().startswith("GB-") else gb_pn
        idx = rest.rfind("-")
        if idx != -1:
            parsed_style = rest[:idx].strip()
            parsed_color = rest[idx+1:].strip()
            # Map abbreviation to full style
            style_map_local = {
                "rec m": "Rectangle M",
                "rec l": "Rectangle L",
                "rec": "Rectangle M",
                "panthos m": "Panthos M",
                "panthos s": "Panthos S",
                "pantos s": "Pantos S",
                "bold": "Bold",
                "slim": "Slim",
                "cateye": "Cateye",
            }
            full_parsed_style = style_map_local.get(parsed_style.lower(), parsed_style)
            # Try to find matching canonical GB by Style+Color
            for canon_gb in sku_gb_lower_to_canonical.values():
                sc = gb_style_color.get(canon_gb)
                if not sc:
                    continue
                # Color match: exact or both low cost
                cmatch = False
                if sc[1].lower() == parsed_color.lower():
                    cmatch = True
                elif "low cost" in parsed_color.lower() and "low cost" in sc[1].lower():
                    cmatch = True
                if not cmatch:
                    continue
                if full_parsed_style.lower() in sc[0].lower() or sc[0].lower() in full_parsed_style.lower():
                    return canon_gb
        return gb_pn

    # Canonicalize daily dicts for GBs (merge case variants and alias)
    def canonicalize_gb_dict(d):
        new_dict = {}
        for pn, daily in d.items():
            if pn.startswith("GB-"):
                canon = to_canonical_gb(pn)
                if canon in new_dict:
                    # Merge daily quantities
                    for ds, qty in daily.items():
                        new_dict[canon][ds] = new_dict[canon].get(ds, 0.0) + float(qty)
                else:
                    new_dict[canon] = dict(daily)
            else:
                # For non-GB (shouldn't happen here) keep as is
                new_dict[pn] = daily
        return new_dict

    # Apply canonicalization to gated, ungated, ctb_gb (ctb_gb only contains GBs)
    # Keep original for SKU daily (they are not GB), but for GB keys we canonicalize
    # For gated/ungated, they contain both SKU and GB; we need to keep SKU keys untouched, only GB keys canonicalized
    # So we handle separately: keep SKU entries as is, canonicalize GB entries and merge
    def canonicalize_mixed_dict(d):
        # d is {PN: daily} where PN can be SKU or GB
        new_dict = {}
        for pn, daily in d.items():
            if pn in sku_attrs:
                # SKU, keep original PN
                new_dict[pn] = daily
            else:
                # GB or other, canonicalize
                canon = to_canonical_gb(pn) if pn.startswith("GB-") else pn
                if canon in new_dict:
                    for ds, qty in daily.items():
                        new_dict[canon][ds] = new_dict[canon].get(ds, 0.0) + float(qty)
                else:
                    new_dict[canon] = dict(daily)
        return new_dict

    plan_gated = canonicalize_mixed_dict(plan_gated)
    plan_ungated = canonicalize_mixed_dict(plan_ungated)
    ctb_gb = canonicalize_gb_dict(ctb_gb)

    # ---- Build GB set: from mapping + any PN in gated/ungated/ctb_gb that is not a SKU ----
    all_gb_set = set()
    for g in sku_to_gb.values():
        if g:
            all_gb_set.add(g)
    for source in (plan_gated, plan_ungated, ctb_gb):
        for pn in source.keys():
            if pn not in sku_attrs:
                all_gb_set.add(pn)

    gb_groups = defaultdict(list)
    for sku in all_skus:
        g = sku_to_gb.get(sku, "")
        if g:
            # Also canonicalize the GB mapping? sku_to_gb values are already canonical SKU master names,
            # but ensure we use canonical
            g_canon = to_canonical_gb(g)
            gb_groups[g_canon].append(sku)
    # Ensure every GB in set has entry even if no SKU maps
    for gb in all_gb_set:
        if gb not in gb_groups:
            gb_groups[gb] = []

    sku_data = {(r["PN"], r["Version-Type"], r["Version-Detail"]): {w: r.get(w) for w in all_weeks} for r in rows}

    # Helper: infer GB Style/Color when exact mapping missing (e.g. Black (low cost) vs DEEP BLACK)
    def _infer_gb_style_color(gb_pn):
        # 1. exact
        if gb_pn in gb_style_color:
            return gb_style_color[gb_pn]
        # 2. case-insensitive
        low = gb_pn.lower()
        for k, v in gb_style_color.items():
            if k.lower() == low:
                return v
        # 3. alias: try to parse GB-Style-Color and find matching SKU by color
        #    e.g. GB-Rec M-Black (low cost) -> style_raw=Rec M, color_raw=Black (low cost)
        #    Should map to DEEP BLACK entries: Color Black (low cost)
        parsed_style = ""
        parsed_color = ""
        rest = gb_pn[3:] if gb_pn.upper().startswith("GB-") else gb_pn
        idx = rest.rfind("-")
        if idx != -1:
            parsed_style = rest[:idx].strip()
            parsed_color = rest[idx+1:].strip()
        else:
            parsed_style = rest.strip()

        # Try to find SKU whose Color matches parsed_color (case-insensitive)
        if parsed_color:
            # For style matching, map abbreviation to full
            style_map_local = {
                "rec m": "Rectangle M",
                "rec l": "Rectangle L",
                "rec": "Rectangle M",
            }
            full_style_local = style_map_local.get(parsed_style.lower(), parsed_style)

            # 1. Exact color match + style match if possible
            # First try with style matching
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                s = attrs.get("Style", "")
                if c.lower() == parsed_color.lower():
                    if full_style_local.lower() in s.lower() or s.lower() in full_style_local.lower() or \
                       parsed_style.lower() in s.lower() or s.lower() in parsed_style.lower():
                        return (s, c)
            # Then any exact color match
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                if c.lower() == parsed_color.lower():
                    return (attrs.get("Style", full_style_local), c)

            # 2. Black (low cost) special: gated uses "Black (low cost)", sku_master uses DEEP BLACK but color is Black (low cost)
            if "low cost" in parsed_color.lower():
                # Find canonical low cost color with style match
                for sku, attrs in sku_attrs.items():
                    c = attrs.get("Color", "")
                    if "low cost" in c.lower():
                        s = attrs.get("Style", "")
                        if full_style_local.lower() in s.lower() or s.lower() in full_style_local.lower():
                            return (s, c)
                # Fallback
                for sku, attrs in sku_attrs.items():
                    c = attrs.get("Color", "")
                    if "low cost" in c.lower():
                        return (full_style_local, c)
            # 3. For "Black Ice" vs "Black Ice (translucent)" - allow contains but require style match
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                s = attrs.get("Style", "")
                # Require style to match as well to avoid BLACK matching Black Ice
                style_match = parsed_style.lower() in s.lower() or s.lower() in parsed_style.lower()
                if not style_match:
                    continue
                if parsed_color.lower() in c.lower() or c.lower() in parsed_color.lower():
                    # e.g. Black Ice vs Black Ice (translucent)
                    return (s, c)

        # Fallback: use parsed as is, try to map style abbreviation to full name
        full_style = parsed_style
        # Map abbreviations: Rec M -> Rectangle M, Rec L -> Rectangle L
        style_map = {
            "rec m": "Rectangle M",
            "rec l": "Rectangle L",
            "rec": "Rectangle M",
            "panthos m": "Panthos M",
            "panthos s": "Panthos S",
            "pantos s": "Pantos S",
            "bold": "Bold",
            "slim": "Slim",
            "cateye": "Cateye",
        }
        low_style = parsed_style.lower()
        if low_style in style_map:
            full_style = style_map[low_style]
        else:
            # Try to find best matching full style from existing sku_attrs
            for sku, attrs in sku_attrs.items():
                s = attrs.get("Style", "")
                if low_style in s.lower() or s.lower() in low_style:
                    full_style = s
                    break

        return (full_style, parsed_color)

    def _infer_gb_usage(gb_pn, skus_list, parsed_color=None):
        # Try from direct SKUs first
        for s in skus_list:
            u = sku_attrs.get(s, {}).get("Usage", "")
            if u:
                return u
        # Try to infer from color/style matching
        if parsed_color is None:
            rest = gb_pn[3:] if gb_pn.upper().startswith("GB-") else gb_pn
            idx = rest.rfind("-")
            parsed_color = rest[idx+1:].strip() if idx != -1 else ""
        if parsed_color:
            # Exact
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                if c.lower() == parsed_color.lower():
                    u = attrs.get("Usage", "")
                    if u:
                        return u
            # Low cost
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                if "low cost" in parsed_color.lower() and "low cost" in c.lower():
                    u = attrs.get("Usage", "")
                    if u:
                        return u
            # Contains (e.g. Black Ice vs Black Ice (translucent))
            for sku, attrs in sku_attrs.items():
                c = attrs.get("Color", "")
                if parsed_color.lower() in c.lower() or c.lower() in parsed_color.lower():
                    u = attrs.get("Usage", "")
                    if u:
                        return u
        # Fallback: MP is default for most
        return "MP"

    for gb, skus in gb_groups.items():
        sc = _infer_gb_style_color(gb)
        # If still blank and skus exist, try first sku's style/color
        if (not sc[0] or not sc[1]) and skus:
            for s in skus:
                attrs = sku_attrs.get(s, {})
                if attrs.get("Style") or attrs.get("Color"):
                    sc = (sc[0] or attrs.get("Style",""), sc[1] or attrs.get("Color",""))
                    if sc[0] and sc[1]:
                        break

        gb_usage = ""
        for s in skus:
            u = sku_attrs.get(s, {}).get("Usage", "")
            if u:
                gb_usage = u
                break
        if not gb_usage:
            gb_usage = _infer_gb_usage(gb, skus)

        base = {"PN": gb, "Usage": gb_usage, "Style": sc[0], "Color": sc[1],
                "GB_PN": gb, "Pallet_Qty": DEFAULT_PALLET_QTY, "_dim": "GB"}

        # ---- ExF: sum SKU ExF (GB has no direct ExF) ----
        exf_vals = {}
        for w in all_weeks:
            s = 0
            has = False
            for sku in skus:
                v = sku_data.get((sku, "ExF", ""), {}).get(w)
                if v is not None:
                    has = True
                    s += v or 0
            if s > 0:
                exf_vals[w] = round(s, 0)
            elif has:
                # keep 0 if needed? but original kept only >0, we keep >0 only for simplicity, but allow 0 diff logic later
                pass

        # ---- Direct daily for GB ----
        gated_daily = plan_gated.get(gb, {})
        ungated_daily = plan_ungated.get(gb, {})

        def get_gb_pack(daily):
            if not daily:
                return {}
            return aggregate_cumulative(daily, cfg["gb_cut"])

        def get_gb_etd(daily):
            if not daily:
                return {}
            raw = aggregate_etd_from_packout(daily, cfg["gb_cut"], offset_n)
            rounded = {}
            for wk, v in raw.items():
                if v and v > 0:
                    rounded[wk] = (int(v) // DEFAULT_PALLET_QTY) * DEFAULT_PALLET_QTY
            return rounded

        # Gated Packout direct, else fallback sum SKU
        if gated_daily:
            gp_vals = get_gb_pack(gated_daily)
        else:
            gp_vals = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Gated", "Packout"), {}).get(w, 0) or 0 for s in skus)
                if s > 0:
                    gp_vals[w] = round(s, 0)

        # Ungated Packout
        if ungated_daily:
            up_vals = get_gb_pack(ungated_daily)
        else:
            up_vals = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Ungated", "Packout"), {}).get(w, 0) or 0 for s in skus)
                if s > 0:
                    up_vals[w] = round(s, 0)

        # Gated ETD
        if gated_daily:
            ge_vals = get_gb_etd(gated_daily)
        else:
            ge_vals = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Gated", "ETD"), {}).get(w, 0) or 0 for s in skus)
                if s > 0:
                    ge_vals[w] = round(s, 0)

        # Ungated ETD
        if ungated_daily:
            ue_vals = get_gb_etd(ungated_daily)
        else:
            ue_vals = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Ungated", "ETD"), {}).get(w, 0) or 0 for s in skus)
                if s > 0:
                    ue_vals[w] = round(s, 0)

        # Diffs: direct GB vs ExF = (base - exf) where either exists
        def direct_vs(base_vals, exf_v):
            res = {}
            for w in all_weeks:
                b = base_vals.get(w)
                e = exf_v.get(w)
                if b is None and e is None:
                    continue
                res[w] = round((b or 0) - (e or 0), 0)
            return res

        gp_vs = direct_vs(gp_vals, exf_vals) if gated_daily else {}
        up_vs = direct_vs(up_vals, exf_vals) if ungated_daily else {}
        ge_vs = direct_vs(ge_vals, exf_vals) if gated_daily else {}
        ue_vs = direct_vs(ue_vals, exf_vals) if ungated_daily else {}

        # Fallback vs logic for GB without direct daily (sum SKU diffs with has_data)
        if not gated_daily:
            gp_vs = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Gated", "Packout vs ExF"), {}).get(w, 0) or 0 for s in skus)
                has_data = False
                for sku in skus:
                    if sku_data.get((sku, "Gated", "Packout vs ExF"), {}).get(w) is not None:
                        has_data = True
                        break
                    if sku_data.get((sku, "Gated", "Packout"), {}).get(w) is not None:
                        has_data = True
                        break
                    if sku_data.get((sku, "ExF", ""), {}).get(w) is not None:
                        has_data = True
                        break
                if has_data:
                    gp_vs[w] = round(s, 0)
        if not ungated_daily:
            up_vs = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Ungated", "Packout vs ExF"), {}).get(w, 0) or 0 for s in skus)
                has_data = False
                for sku in skus:
                    if sku_data.get((sku, "Ungated", "Packout vs ExF"), {}).get(w) is not None:
                        has_data = True
                        break
                    if sku_data.get((sku, "Ungated", "Packout"), {}).get(w) is not None:
                        has_data = True
                        break
                    if sku_data.get((sku, "ExF", ""), {}).get(w) is not None:
                        has_data = True
                        break
                if has_data:
                    up_vs[w] = round(s, 0)
        if not gated_daily:
            ge_vs = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Gated", "ETD vs ExF"), {}).get(w, 0) or 0 for s in skus)
                has_data = False
                for sku in skus:
                    if sku_data.get((sku, "Gated", "ETD vs ExF"), {}).get(w) is not None or \
                       sku_data.get((sku, "Gated", "ETD"), {}).get(w) is not None or \
                       sku_data.get((sku, "ExF", ""), {}).get(w) is not None:
                        has_data = True
                        break
                if has_data:
                    ge_vs[w] = round(s, 0)
        if not ungated_daily:
            ue_vs = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "Ungated", "ETD vs ExF"), {}).get(w, 0) or 0 for s in skus)
                has_data = False
                for sku in skus:
                    if sku_data.get((sku, "Ungated", "ETD vs ExF"), {}).get(w) is not None or \
                       sku_data.get((sku, "Ungated", "ETD"), {}).get(w) is not None or \
                       sku_data.get((sku, "ExF", ""), {}).get(w) is not None:
                        has_data = True
                        break
                if has_data:
                    ue_vs[w] = round(s, 0)

        # CTB
        if gb in ctb_gb:
            ctb_vals_raw = extract_weekly_cum(ctb_gb[gb], cfg["etd_cut"])
            ctb_vals = {w: round(v, 0) for w, v in ctb_vals_raw.items() if v and v > 0}
        else:
            ctb_vals = {}
            for w in all_weeks:
                s = sum(sku_data.get((s, "CTB", ""), {}).get(w, 0) or 0 for s in skus)
                if s > 0:
                    ctb_vals[w] = round(s, 0)

        # Build rows: include ETD as well to match direct availability
        version_defs = [
            ("ExF", "", exf_vals, cfg["exf_cut"]),
            ("Ungated", "ETD", ue_vals, cfg["gb_cut"]),
            ("Ungated", "ETD vs ExF", ue_vs, cfg["gb_cut"]),
            ("Ungated", "Packout", up_vals, cfg["gb_cut"]),
            ("Ungated", "Packout vs ExF", up_vs, cfg["gb_cut"]),
            ("Gated", "ETD", ge_vals, cfg["gb_cut"]),
            ("Gated", "ETD vs ExF", ge_vs, cfg["gb_cut"]),
            ("Gated", "Packout", gp_vals, cfg["gb_cut"]),
            ("Gated", "Packout vs ExF", gp_vs, cfg["gb_cut"]),
            ("CTB", "", ctb_vals, ""),
        ]

        for vt, vd, vals, cd in version_defs:
            rows.append({**base, "Version-Type": vt, "Version-Detail": vd, "Cut Day": cd, **fill(vals)})

    wl = {}
    for w in all_weeks:
        try:
            dt = _to_dt(w)
            wk = (dt.day - 1) // 7 + 1
            wl[w] = f"{dt.strftime('%b')} Wk{wk} ({dt.strftime('%b %d')})"
        except:
            wl[w] = w
    warnings = []
    if missing_sheets:
        warnings.append(f"Missing optional sheets (left empty): {', '.join(missing_sheets)}")
    else:
        warnings.append("All 6 sheets present ✓")

    return {"rows": rows, "weeks": all_weeks, "week_labels": wl, "config": cfg, "warnings": warnings}


def _write_sheet(ws, rows, weeks, fixed, flabels, fills, hf, hfl, hb, cf, cb, nf, wlabels):
    from openpyxl.styles import Alignment as A, Font as F, Border as B, Side as S
    for ci, lab in enumerate(flabels + wlabels, 1):
        c = ws.cell(1, ci, lab)
        c.font = hf; c.fill = hfl; c.border = hb
        c.alignment = A(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28

    last_pn = None
    for ri, r in enumerate(rows, 2):
        vals = []
        for k in fixed:
            if k == "_dim": vals.append(r.get("_dim", ""))
            elif k == "Pallet_Qty": vals.append(r.get("Pallet_Qty") or "")
            else: vals.append(str(r.get(k, "") or ""))
        for w in weeks:
            v = r.get(w)
            vals.append(v if v is not None else "")
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = cf; cell.border = cb
            cell.alignment = A(vertical='center')
            if ci > len(fixed) and isinstance(val, (int, float)):
                cell.number_format = nf
                cell.alignment = A(horizontal='right', vertical='center')
                if val > 0:
                    cell.font = F(name="微软雅黑", size=10, color="059669")
                elif val < 0:
                    cell.font = F(name="微软雅黑", size=10, color="DC2626")
        ft = fills.get(r.get("Version-Type"))
        if ft:
            for ci in range(1, len(fixed) + len(weeks) + 1):
                ws.cell(ri, ci).fill = ft
        is_new = r.get("PN") != last_pn
        if is_new and ri > 2:
            for ci in range(1, len(fixed) + len(weeks) + 1):
                c = ws.cell(ri, ci)
                c.border = B(left=S(style='thin'), right=S(style='thin'),
                             top=S(style='medium', color="94A3B8"), bottom=S(style='thin'))
        last_pn = r.get("PN")
    ws.freeze_panes = ws.cell(2, len(fixed) + 1)


def generate_excel(data):
    """Generate formatted xlsx from report data — split into FG and GB sheets"""
    import openpyxl as xl
    from openpyxl.styles import Font as F, PatternFill as PF, Border as B, Side as S, Alignment as A

    rows, weeks = data["rows"], data["weeks"]
    fg_rows = [r for r in rows if r.get("_dim") == "FG"]
    gb_rows = [r for r in rows if r.get("_dim") == "GB"]

    wb = xl.Workbook()

    hf = F(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    hfl = PF("solid", fgColor="1E293B")
    hb = B(left=S(style='thin'), right=S(style='thin'), top=S(style='thin'), bottom=S(style='thin'))
    cf = F(name="微软雅黑", size=10)
    cb = B(left=S(style='thin'), right=S(style='thin'), top=S(style='thin'), bottom=S(style='thin'))
    nf = '#,##0'
    fills = {"ExF": PF("solid", fgColor="F0F9FF"), "Ungated": PF("solid", fgColor="F0FDF4"),
             "Gated": PF("solid", fgColor="FFFBEB"), "CTB": PF("solid", fgColor="FAF5FF")}

    fixed = ["_dim", "PN", "Usage", "Style", "Color", "Version-Type", "Version-Detail", "Cut Day", "Pallet_Qty"]
    flabels = ["Dim", "PN", "Usage", "Style", "Color", "Version-Type", "Version-Detail", "Cut Day", "Pallet"]

    wlabels = []
    for w in weeks:
        try:
            d = _to_dt(w)
            wk = (d.day - 1) // 7 + 1
            wlabels.append(f"{d.strftime('%b')} Wk{wk} ({d.strftime('%b %d')})")
        except:
            wlabels.append(w)

    ws_fg = wb.active
    ws_fg.title = "FG"
    _write_sheet(ws_fg, fg_rows, weeks, fixed, flabels, fills, hf, hfl, hb, cf, cb, nf, wlabels)

    ws_gb = wb.create_sheet("GB")
    _write_sheet(ws_gb, gb_rows, weeks, fixed, flabels, fills, hf, hfl, hb, cf, cb, nf, wlabels)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
