"""Plan merge module: shared helper functions"""

from datetime import datetime

_DATE_FORMATS = [
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%Y年%m月%d日",
    "%Y年%m月%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%Y%m%d",
]

def normalize_date_str(value):
    """Convert datetime or date-like string to 'YYYY-MM-DD' (or return as-is)."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if not isinstance(value, str):
        return str(value).strip() if value is not None else ""
    s = value.strip()
    if not s:
        return ""
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s


def read_sheet(ws, key_col=None):
    """Read a worksheet → {PN: {date_str: value}}"""
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if key_col is None:
        key_col = "PN" if "PN" in headers else ("SKU" if "SKU" in headers else headers[0])
    ki = headers.index(key_col) if key_col in headers else 0
    result = {}
    for r in range(2, ws.max_row + 1):
        pn = str(ws.cell(r, ki + 1).value or "").strip()
        if not pn: continue
        vals = {}
        for ci, h in enumerate(headers):
            if ci == ki: continue
            v = ws.cell(r, ci + 1).value
            h_str = normalize_date_str(h)
            if h_str and v is not None and isinstance(v, (int, float)):
                vals[h_str] = float(v)
        if vals:
            result[pn] = vals
    return result


def read_sku_master_from_ws(ws):
    """Read sku_master sheet → attributes dict"""
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    sku_attrs, sku_to_gb, sku_pallet, gb_style_color = {}, {}, {}, {}
    for r in range(2, ws.max_row + 1):
        row = {headers[ci]: ws.cell(r, ci + 1).value for ci in range(len(headers))}
        sku = str(row.get("SKU", "") or "").strip()
        if not sku: continue
        # Trim Style/Color/Usage to avoid "Dark Havana " vs "Dark Havana" duplicates in filters
        def _clean(v):
            return str(v).strip() if v is not None else ""
        sku_attrs[sku] = {"Style": _clean(row.get("Style","")),
                           "Color": _clean(row.get("Color","")),
                           "Usage": _clean(row.get("Usage",""))}
        gb = str(row.get("GB_PN","") or "").strip()
        sku_to_gb[sku] = gb
        try:
            pallet_raw = row.get("Pallet_Qty", 864)
            sku_pallet[sku] = int(float(pallet_raw)) if pallet_raw not in (None, "") else 864
        except Exception:
            sku_pallet[sku] = 864
        if gb and row.get("Style") and row.get("Color"):
            gb_style_color[gb] = (_clean(row["Style"]), _clean(row["Color"]))
    return sku_attrs, sku_to_gb, sku_pallet, gb_style_color


def read_uploaded_xlsx(fp):
    """Read single xlsx → (data_dict, missing_optional_sheets_list)"""
    import openpyxl
    wb = openpyxl.load_workbook(fp, data_only=True)
    sheet_names = [s.title for s in wb.worksheets]
    result = {}
    for sn in sheet_names:
        ws = wb[sn]
        if sn == "sku_master":
            result["sku"] = ws
        elif sn == "plan_output_gated":
            result["gated"] = read_sheet(ws)
        elif sn == "plan_output_ungated":
            result["ungated"] = read_sheet(ws)
        elif sn == "forecast":
            result["fcst"] = read_sheet(ws)
        elif sn in ("ctb_sku_cum", "ctb_cum"):
            result["ctb"] = read_sheet(ws, "SKU")
        elif sn in ("ctb_gb_cum", "ctb_gb"):
            result["ctb_gb"] = read_sheet(ws)

    optional_groups = [
        (["plan_output_gated"], "gated"),
        (["plan_output_ungated"], "ungated"),
        (["forecast"], "fcst"),
        (["ctb_sku_cum", "ctb_cum"], "ctb"),
        (["ctb_gb_cum", "ctb_gb"], "ctb_gb"),
    ]
    missing = []
    for names, _key in optional_groups:
        if not any(n in sheet_names for n in names):
            missing.append(names[0])
    return result, missing
